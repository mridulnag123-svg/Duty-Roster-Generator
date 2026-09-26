import os
import sys
import subprocess
import threading
import shutil
import base64
from pathlib import Path
from datetime import datetime

import pandas as pd
import gradio as gr


# ============================================================
# RKM DUTY ROSTER GENERATOR
# STEP 7.5
# FINAL GUI CONTROLS & USABILITY
# ============================================================

APP_TITLE = "RKM Duty Roster Generator"

LOGIN_USERNAME = "admin"
LOGIN_PASSWORD = "rkm123"

BASE_DIR = Path(__file__).resolve().parent

APP_FILE = BASE_DIR / "app.py"

AVAILABILITY_FILE = (
    BASE_DIR
    / "data"
    / "input"
    / "availability.xlsx"
)

STAFF_MASTER_FILE = (
    BASE_DIR
    / "data"
    / "master"
    / "staff_master.xlsx"
)

DUTY_MASTER_FILE = (
    BASE_DIR
    / "data"
    / "master"
    / "duty_master.xlsx"
)

OUTPUT_DIR = (
    BASE_DIR
    / "data"
    / "output"
)

OUTPUT_FILE = (
    OUTPUT_DIR
    / "duty_roster.xlsx"
)

SOURCE_LOGO_FILE = (
    BASE_DIR
    / ".venv"
    / "Lib"
    / "site-packages"
    / "assets"
    / "logo.png"
)

RKM_LOGO_FILE = (
    BASE_DIR
    / "public"
    / "logo.png"
)


def ensure_logo_asset():
    if SOURCE_LOGO_FILE.exists() and not RKM_LOGO_FILE.exists():
        RKM_LOGO_FILE.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(SOURCE_LOGO_FILE, RKM_LOGO_FILE)
    return str(RKM_LOGO_FILE) if RKM_LOGO_FILE.exists() else ""


# ============================================================
# GLOBAL STATE
# ============================================================

original_roster_df = pd.DataFrame()

generation_lock = threading.Lock()

generation_running = False


# ============================================================
# REQUIRED COLUMNS
# ============================================================

AVAILABILITY_REQUIRED_COLUMNS = [
    "Date",
    "Available",
    "Leave",
]

STAFF_REQUIRED_COLUMNS = [
    "Staff Name",
    "Active",
]

DUTY_REQUIRED_COLUMNS = [
    "Duty",
    "Required Staff",
]

# Standard RKM duties used by the dynamic 3/4/5-staff duty structure.
STANDARD_RKM_DUTIES = [
    "Morning",
    "Afternoon",
    "Hostel Afternoon",
    "Hostel Night",
    "Night",
]


# ============================================================
# SAFE EMPTY DATAFRAME
# ============================================================

def empty_df():
    return pd.DataFrame()


def login_user(username, password):
    if username == LOGIN_USERNAME and password == LOGIN_PASSWORD:
        return gr.update(visible=False), "", gr.update(visible=True)

    return gr.update(visible=True), "Invalid username or password.", gr.update(visible=False)


def logout_user():
    return gr.update(visible=True), gr.update(visible=False)


# ============================================================
# SAFE READ EXCEL
# ============================================================

def safe_read_excel(
    file_path,
    sheet_name=0
):

    try:

        file_path = Path(file_path)

        if not file_path.exists():

            return (
                None,
                f"File not found: {file_path}"
            )

        df = pd.read_excel(
            file_path,
            sheet_name=sheet_name
        )

        if df is None:

            df = pd.DataFrame()

        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
        )

        return (
            df,
            None
        )

    except Exception as e:

        return (
            None,
            str(e)
        )


# ============================================================
# FILE STATUS
# ============================================================

def get_file_status():

    files = [
        (
            "Availability",
            AVAILABILITY_FILE
        ),
        (
            "Staff Master",
            STAFF_MASTER_FILE
        ),
        (
            "Duty Master",
            DUTY_MASTER_FILE
        ),
        (
            "app.py",
            APP_FILE
        ),
    ]

    lines = []

    all_exist = True

    for label, path in files:

        if path.exists():

            size_kb = (
                path.stat().st_size
                / 1024
            )

            modified = (
                datetime.fromtimestamp(
                    path.stat().st_mtime
                )
                .strftime(
                    "%d/%m/%Y %H:%M"
                )
            )

            lines.append(
                f"🟢 **{label}: READY**  \n"
                f"`{path.name}` | "
                f"{size_kb:.1f} KB | "
                f"Modified: {modified}"
            )

        else:

            all_exist = False

            lines.append(
                f"🔴 **{label}: MISSING**  \n"
                f"`{path}`"
            )

    if OUTPUT_FILE.exists():

        size_kb = (
            OUTPUT_FILE.stat().st_size
            / 1024
        )

        modified = (
            datetime.fromtimestamp(
                OUTPUT_FILE.stat().st_mtime
            )
            .strftime(
                "%d/%m/%Y %H:%M"
            )
        )

        lines.append(
            f"📘 **Previous Output: AVAILABLE**  \n"
            f"`{OUTPUT_FILE.name}` | "
            f"{size_kb:.1f} KB | "
            f"Modified: {modified}"
        )

    else:

        lines.append(
            "⚪ **Previous Output: NOT GENERATED YET**"
        )

    if all_exist:

        title = (
            "### 🟢 INPUT FILE STATUS: READY"
        )

    else:

        title = (
            "### 🔴 INPUT FILE STATUS: INCOMPLETE"
        )

    return (
        title
        + "\n\n"
        + "\n\n".join(lines)
    )


# ============================================================
# OPEN FOLDER
# ============================================================

def open_folder(folder_path):

    try:

        folder_path = Path(
            folder_path
        )

        folder_path.mkdir(
            parents=True,
            exist_ok=True
        )

        if os.name == "nt":

            os.startfile(
                str(folder_path)
            )

        elif sys.platform == "darwin":

            subprocess.Popen(
                [
                    "open",
                    str(folder_path)
                ]
            )

        else:

            subprocess.Popen(
                [
                    "xdg-open",
                    str(folder_path)
                ]
            )

        return (
            "### 🟢 Folder opened successfully\n\n"
            f"`{folder_path}`"
        )

    except Exception as e:

        return (
            "### 🔴 Could not open folder\n\n"
            f"`{e}`"
        )


def open_input_folder():

    return open_folder(
        AVAILABILITY_FILE.parent
    )


def open_output_folder():

    return open_folder(
        OUTPUT_DIR
    )


# ============================================================
# VALIDATE ONE FILE
# ============================================================

def validate_file(
    label,
    file_path,
    required_columns
):

    df, error = safe_read_excel(
        file_path
    )

    if error:

        return {
            "File": label,
            "Status": "FAIL",
            "Details": error,
            "Rows": 0,
        }

    missing = [
        col
        for col in required_columns
        if col not in df.columns
    ]

    if missing:

        return {
            "File": label,
            "Status": "FAIL",
            "Details": (
                "Missing columns: "
                + ", ".join(missing)
            ),
            "Rows": len(df),
        }

    return {
        "File": label,
        "Status": "PASS",
        "Details": (
            "Required columns found."
        ),
        "Rows": len(df),
    }


# ============================================================
# CHECK INPUT FILES
# ============================================================

def check_input_files():

    results = []

    results.append(
        validate_file(
            "Availability",
            AVAILABILITY_FILE,
            AVAILABILITY_REQUIRED_COLUMNS
        )
    )

    results.append(
        validate_file(
            "Staff Master",
            STAFF_MASTER_FILE,
            STAFF_REQUIRED_COLUMNS
        )
    )

    results.append(
        validate_file(
            "Duty Master",
            DUTY_MASTER_FILE,
            DUTY_REQUIRED_COLUMNS
        )
    )

    if APP_FILE.exists():

        results.append({
            "File": "app.py",
            "Status": "PASS",
            "Details": "Roster engine found.",
            "Rows": 0,
        })

    else:

        results.append({
            "File": "app.py",
            "Status": "FAIL",
            "Details": "app.py not found.",
            "Rows": 0,
        })

    result_df = pd.DataFrame(
        results
    )

    pass_count = int(
        result_df["Status"]
        .eq("PASS")
        .sum()
    )

    fail_count = int(
        result_df["Status"]
        .eq("FAIL")
        .sum()
    )

    if fail_count == 0:

        message = (
            "### 🟢 INPUT VALIDATION PASSED\n\n"
            f"All **{pass_count}** required "
            "components are ready."
        )

    else:

        message = (
            "### 🔴 INPUT VALIDATION FAILED\n\n"
            f"PASS: **{pass_count}**  |  "
            f"FAIL: **{fail_count}**\n\n"
            "Please correct the failed item(s) "
            "before generating the roster."
        )

    return (
        message,
        result_df
    )


# ============================================================
# LOAD GENERATED ROSTER
# ============================================================

def load_generated_roster():

    if not OUTPUT_FILE.exists():

        return (
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            "No output file found."
        )

    try:

        excel = pd.ExcelFile(
            OUTPUT_FILE
        )

        sheets = excel.sheet_names

        # ----------------------------------------------------
        # ROSTER
        # ----------------------------------------------------

        if "Roster" in sheets:

            roster_df = pd.read_excel(
                OUTPUT_FILE,
                sheet_name="Roster"
            )

        else:

            roster_df = pd.DataFrame()

        # ----------------------------------------------------
        # STAFF SUMMARY
        # ----------------------------------------------------

        if "Staff Summary" in sheets:

            staff_summary_df = pd.read_excel(
                OUTPUT_FILE,
                sheet_name="Staff Summary"
            )

        else:

            staff_summary_df = pd.DataFrame()

        # ----------------------------------------------------
        # DUTY SUMMARY
        # ----------------------------------------------------

        if "Duty Summary" in sheets:

            duty_summary_df = pd.read_excel(
                OUTPUT_FILE,
                sheet_name="Duty Summary"
            )

        else:

            duty_summary_df = pd.DataFrame()

        # ----------------------------------------------------
        # VALIDATION
        # ----------------------------------------------------

        if "Validation" in sheets:

            validation_df = pd.read_excel(
                OUTPUT_FILE,
                sheet_name="Validation"
            )

        else:

            validation_df = pd.DataFrame()

        # ----------------------------------------------------
        # COVER REPORT
        # ----------------------------------------------------

        if "Cover Report" in sheets:

            cover_df = pd.read_excel(
                OUTPUT_FILE,
                sheet_name="Cover Report"
            )

        else:

            cover_df = pd.DataFrame()

        return (
            roster_df,
            staff_summary_df,
            duty_summary_df,
            validation_df,
            cover_df,
            "Output loaded successfully."
        )

    except Exception as e:

        return (
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            f"Could not read output: {e}"
        )


# ============================================================
# NORMALIZE ROSTER
# ============================================================

def normalize_roster(df):

    if df is None or df.empty:

        return pd.DataFrame()

    result = df.copy()

    result.columns = (
        result.columns
        .astype(str)
        .str.strip()
    )

    # --------------------------------------------------------
    # DATE
    # --------------------------------------------------------

    if "Date" in result.columns:

        result["Date"] = pd.to_datetime(
            result["Date"],
            errors="coerce",
            dayfirst=True
        )

    # --------------------------------------------------------
    # TEXT COLUMNS
    # --------------------------------------------------------

    for col in [
        "Duty",
        "Staff",
        "Assignment Type",
        "Reason",
    ]:

        if col in result.columns:

            result[col] = (
                result[col]
                .fillna("")
                .astype(str)
                .str.strip()
            )

    return result


# ============================================================
# DISPLAY ROSTER
# ============================================================

def prepare_display_roster(df):

    result = normalize_roster(
        df
    )

    if result.empty:

        return pd.DataFrame()

    display_df = result.copy()

    if "Date" in display_df.columns:

        display_df["Date"] = (
            display_df["Date"]
            .dt.strftime(
                "%d/%m/%Y"
            )
        )

    return display_df


# ============================================================
# GET FILTER VALUES
# ============================================================

def get_filter_values(
    roster_df
):

    df = normalize_roster(
        roster_df
    )

    dates = ["ALL"]
    staff = ["ALL"]
    duties = ["ALL"]

    if df.empty:

        return (
            dates,
            staff,
            duties
        )

    # --------------------------------------------------------
    # DATES
    # --------------------------------------------------------

    if "Date" in df.columns:

        valid_dates = (
            df["Date"]
            .dropna()
            .dt.strftime(
                "%d/%m/%Y"
            )
            .drop_duplicates()
            .tolist()
        )

        dates.extend(
            sorted(valid_dates)
        )

    # --------------------------------------------------------
    # STAFF
    # --------------------------------------------------------

    if "Staff" in df.columns:

        values = (
            df["Staff"]
            .replace("", pd.NA)
            .dropna()
            .drop_duplicates()
            .tolist()
        )

        staff.extend(
            sorted(values)
        )

    # --------------------------------------------------------
    # DUTY
    # --------------------------------------------------------

    if "Duty" in df.columns:

        values = (
            df["Duty"]
            .replace("", pd.NA)
            .dropna()
            .drop_duplicates()
            .tolist()
        )

        duties.extend(
            sorted(values)
        )

    return (
        dates,
        staff,
        duties
    )


# ============================================================
# FILTER ROSTER
# ============================================================

def apply_roster_filter(
    roster_df,
    selected_date,
    selected_staff,
    selected_duty,
    search_text
):

    df = normalize_roster(
        roster_df
    )

    if df.empty:

        return (
            pd.DataFrame(),
            "### ⚪ No roster data available."
        )

    # --------------------------------------------------------
    # DATE
    # --------------------------------------------------------

    if (
        selected_date
        and selected_date != "ALL"
        and "Date" in df.columns
    ):

        df = df[
            df["Date"]
            .dt.strftime("%d/%m/%Y")
            == selected_date
        ]

    # --------------------------------------------------------
    # STAFF
    # --------------------------------------------------------

    if (
        selected_staff
        and selected_staff != "ALL"
        and "Staff" in df.columns
    ):

        df = df[
            df["Staff"]
            == selected_staff
        ]

    # --------------------------------------------------------
    # DUTY
    # --------------------------------------------------------

    if (
        selected_duty
        and selected_duty != "ALL"
        and "Duty" in df.columns
    ):

        df = df[
            df["Duty"]
            == selected_duty
        ]

    # --------------------------------------------------------
    # SEARCH
    # --------------------------------------------------------

    if search_text:

        search_value = str(
            search_text
        ).strip().lower()

        if search_value:

            mask = pd.Series(
                False,
                index=df.index
            )

            for col in df.columns:

                mask = (
                    mask
                    |
                    df[col]
                    .astype(str)
                    .str.lower()
                    .str.contains(
                        search_value,
                        na=False,
                        regex=False
                    )
                )

            df = df[mask]

    # --------------------------------------------------------
    # DISPLAY
    # --------------------------------------------------------

    display_df = prepare_display_roster(
        df
    )

    # --------------------------------------------------------
    # STATUS
    # --------------------------------------------------------

    if display_df.empty:

        status = (
            "### 🟡 No matching roster "
            "records found.\n\n"
            "Try changing the filters "
            "or search text."
        )

    else:

        status = (
            "### 🟢 Filtered result count: "
            f"**{len(display_df)}** assignment(s)."
        )

    return (
        display_df,
        status
    )


# ============================================================
# CREATE DASHBOARD
# ============================================================

def create_dashboard(
    roster_df,
    staff_summary_df,
    duty_summary_df
):

    df = normalize_roster(
        roster_df
    )

    if df.empty:

        return """
## 📊 Roster Dashboard

| Metric | Value |
|---|---:|
| 📋 Total Assignments | **0** |
| 🟢 PRIMARY | **0** |
| 🟡 COVER | **0** |
| 🔴 UNASSIGNED | **0** |
| 👥 Active Staff | **0** |
| 📅 Dates | **0** |

### ⚪ No roster generated yet
"""

    total = len(df)

    primary = 0
    cover = 0
    unassigned = 0

    if "Assignment Type" in df.columns:

        assignment_types = (
            df["Assignment Type"]
            .astype(str)
            .str.upper()
            .str.strip()
        )

        primary = int(
            (assignment_types == "PRIMARY")
            .sum()
        )

        cover = int(
            (assignment_types == "COVER")
            .sum()
        )

        unassigned = int(
            (assignment_types == "UNASSIGNED")
            .sum()
        )

    # --------------------------------------------------------
    # STAFF
    # --------------------------------------------------------

    if (
        not staff_summary_df.empty
        and "Staff" in staff_summary_df.columns
    ):

        active_staff = (
            staff_summary_df["Staff"]
            .replace("", pd.NA)
            .dropna()
            .nunique()
        )

    elif "Staff" in df.columns:

        active_staff = (
            df["Staff"]
            .replace("", pd.NA)
            .dropna()
            .nunique()
        )

    else:

        active_staff = 0

    # --------------------------------------------------------
    # DATES
    # --------------------------------------------------------

    if "Date" in df.columns:

        total_dates = (
            df["Date"]
            .dropna()
            .nunique()
        )

    else:

        total_dates = 0

    overall = (
        "🔴 WARNING"
        if unassigned > 0
        else "🟢 READY"
    )

    # --------------------------------------------------------
    # DUTY DISTRIBUTION
    # --------------------------------------------------------

    duty_text = ""

    if (
        "Duty" in df.columns
        and not df.empty
    ):

        duty_counts = (
            df["Duty"]
            .value_counts()
        )

        duty_rows = []

        for duty, count in duty_counts.items():

            duty_rows.append(
                f"| {duty} | **{count}** |"
            )

        if duty_rows:

            duty_text = (
                "\n\n"
                "### 📌 Duty Distribution\n\n"
                "| Duty | Assignments |\n"
                "|---|---:|\n"
                + "\n".join(duty_rows)
            )

    return f"""
## {overall} ROSTER DASHBOARD

| Metric | Value |
|---|---:|
| 📋 Total Assignments | **{total}** |
| 🟢 PRIMARY | **{primary}** |
| 🟡 COVER | **{cover}** |
| 🔴 UNASSIGNED | **{unassigned}** |
| 👥 Active Staff | **{active_staff}** |
| 📅 Total Dates | **{total_dates}** |
{duty_text}
"""


# ============================================================
# CREATE SUMMARY
# ============================================================

def create_summary(
    roster_df,
    staff_summary_df,
    duty_summary_df
):

    df = normalize_roster(
        roster_df
    )

    if df.empty:

        return """
## 📊 Generation Summary

### ⚪ No roster generated yet

Click **GENERATE DUTY ROSTER**
to create the roster.
"""

    total = len(df)

    primary = 0
    cover = 0
    unassigned = 0

    if "Assignment Type" in df.columns:

        assignment_types = (
            df["Assignment Type"]
            .astype(str)
            .str.upper()
            .str.strip()
        )

        primary = int(
            (assignment_types == "PRIMARY")
            .sum()
        )

        cover = int(
            (assignment_types == "COVER")
            .sum()
        )

        unassigned = int(
            (assignment_types == "UNASSIGNED")
            .sum()
        )

    if "Date" in df.columns:

        dates = int(
            df["Date"]
            .dropna()
            .nunique()
        )

    else:

        dates = 0

    if "Staff" in df.columns:

        staff = int(
            df["Staff"]
            .replace("", pd.NA)
            .dropna()
            .nunique()
        )

    else:

        staff = 0

    status = (
        "🔴 WARNING"
        if unassigned > 0
        else "🟢 SUCCESS"
    )

    return f"""
## {status} ROSTER GENERATION SUMMARY

| Metric | Result |
|---|---:|
| 📊 Total Assignments | **{total}** |
| 🟢 PRIMARY | **{primary}** |
| 🟡 COVER | **{cover}** |
| 🔴 UNASSIGNED | **{unassigned}** |
| 👥 Staff Used | **{staff}** |
| 📅 Dates | **{dates}** |
| 📁 Output | `duty_roster.xlsx` |

The generated workbook contains
the roster and validation information.
"""


# ============================================================
# GENERATION STATUS
# ============================================================

def get_generation_status():

    if generation_running:

        return (
            "### 🟡 GENERATION IN PROGRESS\n\n"
            "Please wait. A roster is currently "
            "being generated."
        )

    if OUTPUT_FILE.exists():

        modified = (
            datetime.fromtimestamp(
                OUTPUT_FILE.stat().st_mtime
            )
            .strftime(
                "%d/%m/%Y %I:%M:%S %p"
            )
        )

        return (
            "### 🟢 ROSTER AVAILABLE\n\n"
            f"Last generated/updated: **{modified}**"
        )

    return (
        "### ⚪ READY TO GENERATE\n\n"
        "No current roster output found."
    )


# ============================================================
# GENERATION ERROR RESULT
# ============================================================

def generation_error_result(
    status_message,
    log_message=""
):

    return (
        status_message,
        create_dashboard(
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame()
        ),
        create_summary(
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame()
        ),
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        log_message,
        None,
        gr.update(
            choices=["ALL"],
            value="ALL"
        ),
        gr.update(
            choices=["ALL"],
            value="ALL"
        ),
        gr.update(
            choices=["ALL"],
            value="ALL"
        ),
        "### 🔴 No roster available."
    )


# ============================================================
# GENERATE ROSTER
# ============================================================

def generate_roster():

    global generation_running
    global original_roster_df

    # --------------------------------------------------------
    # DOUBLE GENERATION PROTECTION
    # --------------------------------------------------------

    if generation_running:

        return generation_error_result(
            "### 🟡 GENERATION ALREADY IN PROGRESS\n\n"
            "Please wait until the current generation finishes.",
            "Generation already running."
        )

    if not generation_lock.acquire(
        blocking=False
    ):

        return generation_error_result(
            "### 🟡 GENERATION ALREADY IN PROGRESS",
            "Another generation process is active."
        )

    generation_running = True

    try:

        start_time = datetime.now()

        # ----------------------------------------------------
        # VALIDATE INPUT
        # ----------------------------------------------------

        validation_message, validation_df = (
            check_input_files()
        )

        failed = validation_df[
            validation_df["Status"] == "FAIL"
        ]

        if not failed.empty:

            return (
                "### 🔴 GENERATION STOPPED\n\n"
                "Please fix the input validation "
                "errors first.",
                create_dashboard(
                    pd.DataFrame(),
                    pd.DataFrame(),
                    pd.DataFrame()
                ),
                create_summary(
                    pd.DataFrame(),
                    pd.DataFrame(),
                    pd.DataFrame()
                ),
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
                validation_df,
                pd.DataFrame(),
                validation_message,
                None,
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                "### 🔴 No roster available."
            )

        # ----------------------------------------------------
        # OUTPUT DIRECTORY
        # ----------------------------------------------------

        OUTPUT_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        # ----------------------------------------------------
        # OLD OUTPUT HANDLING
        # ----------------------------------------------------
        # Do NOT delete the existing Excel file here.
        # app.py is responsible for creating/replacing the output.
        # On Windows, unlink() can fail because Excel, Explorer
        # preview, antivirus, or another process may temporarily
        # hold the file even when Excel appears closed.
        # We therefore let app.py handle the output file directly.

        if OUTPUT_FILE.exists():

            old_output_note = (
                "\n\n[INFO] Existing duty_roster.xlsx found. "
                "app.py will replace/update the output."
            )

        else:

            old_output_note = ""

        # RUN APP.PY
        # ----------------------------------------------------

        try:

            process = subprocess.run(
                [
                    sys.executable,
                    str(APP_FILE)
                ],
                cwd=str(BASE_DIR),
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace"
            )

            stdout = process.stdout or ""

            stderr = process.stderr or ""

            full_log = stdout

            if stderr.strip():

                full_log += (
                    "\n\n"
                    "========== STDERR ==========\n"
                    + stderr
                )

            full_log += old_output_note

        except Exception as e:

            return generation_error_result(
                "### 🔴 COULD NOT START ROSTER ENGINE.",
                str(e)
            )

        # ----------------------------------------------------
        # PROCESS ERROR
        # ----------------------------------------------------

        if process.returncode != 0:

            error_message = (
                "### 🔴 ROSTER GENERATION FAILED\n\n"
                f"Process exit code: "
                f"**{process.returncode}**"
            )

            if stderr.strip():

                error_message += (
                    "\n\n"
                    "### Error details\n"
                    "```text\n"
                    + stderr.strip()
                    + "\n```"
                )

            return (
                error_message,
                create_dashboard(
                    pd.DataFrame(),
                    pd.DataFrame(),
                    pd.DataFrame()
                ),
                create_summary(
                    pd.DataFrame(),
                    pd.DataFrame(),
                    pd.DataFrame()
                ),
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
                full_log,
                None,
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                "### 🔴 No roster available."
            )

        # ----------------------------------------------------
        # CHECK OUTPUT
        # ----------------------------------------------------

        if not OUTPUT_FILE.exists():

            return (
                "### 🔴 GENERATION FINISHED BUT "
                "OUTPUT FILE IS MISSING\n\n"
                f"Expected file:\n`{OUTPUT_FILE}`",
                create_dashboard(
                    pd.DataFrame(),
                    pd.DataFrame(),
                    pd.DataFrame()
                ),
                create_summary(
                    pd.DataFrame(),
                    pd.DataFrame(),
                    pd.DataFrame()
                ),
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
                full_log,
                None,
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                gr.update(
                    choices=["ALL"],
                    value="ALL"
                ),
                "### 🔴 No roster available."
            )

        # ----------------------------------------------------
        # LOAD GENERATED DATA
        # ----------------------------------------------------

        (
            roster_df,
            staff_summary_df,
            duty_summary_df,
            validation_result_df,
            cover_df,
            load_message
        ) = load_generated_roster()

        # ----------------------------------------------------
        # STORE COMPLETE ORIGINAL ROSTER
        # ----------------------------------------------------

        original_roster_df = (
            roster_df.copy()
        )

        # ----------------------------------------------------
        # FILTER VALUES
        # ----------------------------------------------------

        dates, staff, duties = (
            get_filter_values(
                original_roster_df
            )
        )

        # ----------------------------------------------------
        # DISPLAY
        # ----------------------------------------------------

        display_roster = (
            prepare_display_roster(
                original_roster_df
            )
        )

        # ----------------------------------------------------
        # DASHBOARD
        # ----------------------------------------------------

        dashboard = create_dashboard(
            original_roster_df,
            staff_summary_df,
            duty_summary_df
        )

        # ----------------------------------------------------
        # SUMMARY
        # ----------------------------------------------------

        summary = create_summary(
            original_roster_df,
            staff_summary_df,
            duty_summary_df
        )

        # ----------------------------------------------------
        # TIME
        # ----------------------------------------------------

        elapsed = (
            datetime.now()
            - start_time
        ).total_seconds()

        # ----------------------------------------------------
        # SUCCESS STATUS
        # ----------------------------------------------------

        status = (
            "### 🟢 ROSTER GENERATED SUCCESSFULLY\n\n"
            f"Completed in **{elapsed:.2f} seconds**.\n\n"
            "Generated at: "
            f"**{datetime.now().strftime('%d/%m/%Y %I:%M:%S %p')}**"
        )

        # ----------------------------------------------------
        # FILTER STATUS
        # ----------------------------------------------------

        if display_roster.empty:

            filter_status = (
                "### 🟡 No roster data available."
            )

        else:

            filter_status = (
                "### 🟢 Filtered result count: "
                f"**{len(display_roster)}**"
            )

        # ----------------------------------------------------
        # DOWNLOAD
        # ----------------------------------------------------

        download_value = (
            str(OUTPUT_FILE)
            if OUTPUT_FILE.exists()
            else None
        )

        # IMPORTANT:
        # EXACTLY 14 OUTPUTS
        #
        # 1  generation_status
        # 2  dashboard
        # 3  summary
        # 4  roster_preview
        # 5  staff_summary
        # 6  duty_summary
        # 7  validation
        # 8  cover
        # 9  generation_log
        # 10 download_file
        # 11 date_filter
        # 12 staff_filter
        # 13 duty_filter
        # 14 filter_status

        return (
            status,
            dashboard,
            summary,
            display_roster,
            staff_summary_df,
            duty_summary_df,
            validation_result_df,
            cover_df,
            full_log,
            download_value,
            gr.update(
                choices=dates,
                value="ALL"
            ),
            gr.update(
                choices=staff,
                value="ALL"
            ),
            gr.update(
                choices=duties,
                value="ALL"
            ),
            filter_status
        )

    except Exception as e:

        return generation_error_result(
            "### 🔴 UNEXPECTED GENERATION ERROR\n\n"
            f"```text\n{e}\n```",
            str(e)
        )

    finally:

        generation_running = False

        try:

            generation_lock.release()

        except RuntimeError:

            pass


# ============================================================
# REFRESH EVERYTHING
# ============================================================

def refresh_everything():

    global original_roster_df

    try:

        # ----------------------------------------------------
        # FILE STATUS
        # ----------------------------------------------------

        file_status = get_file_status()

        # ----------------------------------------------------
        # VALIDATION
        # ----------------------------------------------------

        validation_message, validation_df = (
            check_input_files()
        )

        # ----------------------------------------------------
        # LOAD OUTPUT
        # ----------------------------------------------------

        (
            roster_df,
            staff_summary_df,
            duty_summary_df,
            validation_result_df,
            cover_df,
            message
        ) = load_generated_roster()

        # ----------------------------------------------------
        # STORE ORIGINAL
        # ----------------------------------------------------

        original_roster_df = (
            roster_df.copy()
        )

        # ----------------------------------------------------
        # FILTER VALUES
        # ----------------------------------------------------

        dates, staff, duties = (
            get_filter_values(
                original_roster_df
            )
        )

        # ----------------------------------------------------
        # DISPLAY
        # ----------------------------------------------------

        display_roster = (
            prepare_display_roster(
                original_roster_df
            )
        )

        # ----------------------------------------------------
        # DASHBOARD
        # ----------------------------------------------------

        dashboard = create_dashboard(
            original_roster_df,
            staff_summary_df,
            duty_summary_df
        )

        # ----------------------------------------------------
        # SUMMARY
        # ----------------------------------------------------

        summary = create_summary(
            original_roster_df,
            staff_summary_df,
            duty_summary_df
        )

        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        if original_roster_df.empty:

            generation_status = (
                "### ⚪ NO ROSTER GENERATED YET\n\n"
                "Input files can still be checked."
            )

            filter_status = (
                "### ⚪ No roster data available."
            )

        else:

            generation_status = (
                "### 🟢 ROSTER AVAILABLE\n\n"
                f"{len(original_roster_df)} "
                "total assignment(s) loaded."
            )

            filter_status = (
                "### 🟢 Filtered result count: "
                f"**{len(display_roster)}**"
            )

        download_value = (
            str(OUTPUT_FILE)
            if OUTPUT_FILE.exists()
            else None
        )

        return (
            file_status,
            validation_message,
            validation_df,
            generation_status,
            dashboard,
            summary,
            display_roster,
            staff_summary_df,
            duty_summary_df,
            validation_result_df,
            cover_df,
            gr.update(
                choices=dates,
                value="ALL"
            ),
            gr.update(
                choices=staff,
                value="ALL"
            ),
            gr.update(
                choices=duties,
                value="ALL"
            ),
            filter_status,
            download_value
        )

    except Exception as e:

        empty_dashboard = create_dashboard(
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame()
        )

        empty_summary = create_summary(
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame()
        )

        return (
            "### 🔴 REFRESH ERROR\n\n"
            f"`{e}`",
            "### 🔴 Could not refresh validation.",
            pd.DataFrame(),
            "### 🔴 Refresh failed.",
            empty_dashboard,
            empty_summary,
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            gr.update(
                choices=["ALL"],
                value="ALL"
            ),
            gr.update(
                choices=["ALL"],
                value="ALL"
            ),
            gr.update(
                choices=["ALL"],
                value="ALL"
            ),
            "### 🔴 No roster available.",
            None
        )


# ============================================================
# LOAD EXISTING OUTPUT
# ============================================================

def refresh_output():

    global original_roster_df

    (
        roster_df,
        staff_summary_df,
        duty_summary_df,
        validation_df,
        cover_df,
        message
    ) = load_generated_roster()

    original_roster_df = (
        roster_df.copy()
    )

    if original_roster_df.empty:

        return (
            "### ⚪ NO GENERATED ROSTER FOUND.",
            create_dashboard(
                roster_df,
                staff_summary_df,
                duty_summary_df
            ),
            create_summary(
                roster_df,
                staff_summary_df,
                duty_summary_df
            ),
            pd.DataFrame(),
            staff_summary_df,
            duty_summary_df,
            validation_df,
            cover_df,
            gr.update(
                choices=["ALL"],
                value="ALL"
            ),
            gr.update(
                choices=["ALL"],
                value="ALL"
            ),
            gr.update(
                choices=["ALL"],
                value="ALL"
            ),
            pd.DataFrame(),
            "### ⚪ No roster data available.",
            (
                str(OUTPUT_FILE)
                if OUTPUT_FILE.exists()
                else None
            )
        )

    dates, staff, duties = (
        get_filter_values(
            original_roster_df
        )
    )

    display_roster = (
        prepare_display_roster(
            original_roster_df
        )
    )

    return (
        "### 🟢 EXISTING ROSTER LOADED SUCCESSFULLY.",

        create_dashboard(
            original_roster_df,
            staff_summary_df,
            duty_summary_df
        ),

        create_summary(
            original_roster_df,
            staff_summary_df,
            duty_summary_df
        ),

        display_roster,

        staff_summary_df,

        duty_summary_df,

        validation_df,

        cover_df,

        gr.update(
            choices=dates,
            value="ALL"
        ),

        gr.update(
            choices=staff,
            value="ALL"
        ),

        gr.update(
            choices=duties,
            value="ALL"
        ),

        display_roster,

        (
            "### 🟢 Filtered result count: "
            f"**{len(display_roster)}**"
        ),

        (
            str(OUTPUT_FILE)
            if OUTPUT_FILE.exists()
            else None
        )
    )


# ============================================================
# APPLY FILTER BUTTON
# ============================================================

def apply_filter_button(
    selected_date,
    selected_staff,
    selected_duty,
    search_text
):

    global original_roster_df

    return apply_roster_filter(
        original_roster_df,
        selected_date,
        selected_staff,
        selected_duty,
        search_text
    )


# ============================================================
# RESET FILTERS
# ============================================================

def reset_filters():

    global original_roster_df

    if (
        original_roster_df is None
        or original_roster_df.empty
    ):

        return (
            gr.update(
                value="ALL"
            ),
            gr.update(
                value="ALL"
            ),
            gr.update(
                value=""
            ),
            gr.update(
                value="ALL"
            ),
            pd.DataFrame(),
            "### ⚪ No roster data available."
        )

    display_df = (
        prepare_display_roster(
            original_roster_df
        )
    )

    return (
        gr.update(
            value="ALL"
        ),
        gr.update(
            value="ALL"
        ),
        gr.update(
            value=""
        ),
        gr.update(
            value="ALL"
        ),
        display_df,
        (
            "### 🟢 Filtered result count: "
            f"**{len(display_df)}**"
        )
    )


# ============================================================
# EXPORT FILTERED EXCEL (PROFESSIONAL FORMATTING)
# ============================================================

def export_filtered_excel(
    selected_date,
    selected_staff,
    selected_duty,
    search_text
):

    global original_roster_df

    if (
        original_roster_df is None
        or original_roster_df.empty
    ):

        return (
            None,
            "### 🔴 No roster data available for export."
        )

    try:

        df = normalize_roster(
            original_roster_df
        )

        # ----------------------------------------------------
        # DATE
        # ----------------------------------------------------

        if (
            selected_date
            and selected_date != "ALL"
            and "Date" in df.columns
        ):

            df = df[
                df["Date"]
                .dt.strftime("%d/%m/%Y")
                == selected_date
            ]

        # ----------------------------------------------------
        # STAFF
        # ----------------------------------------------------

        if (
            selected_staff
            and selected_staff != "ALL"
            and "Staff" in df.columns
        ):

            df = df[
                df["Staff"]
                == selected_staff
            ]

        # ----------------------------------------------------
        # DUTY
        # ----------------------------------------------------

        if (
            selected_duty
            and selected_duty != "ALL"
            and "Duty" in df.columns
        ):

            df = df[
                df["Duty"]
                == selected_duty
            ]

        # ----------------------------------------------------
        # SEARCH
        # ----------------------------------------------------

        if search_text:

            search_value = str(
                search_text
            ).strip().lower()

            if search_value:

                mask = pd.Series(
                    False,
                    index=df.index
                )

                for col in df.columns:

                    mask = (
                        mask
                        |
                        df[col]
                        .astype(str)
                        .str.lower()
                        .str.contains(
                            search_value,
                            na=False,
                            regex=False
                        )
                    )

                df = df[mask]

        # ----------------------------------------------------
        # EMPTY
        # ----------------------------------------------------

        if df.empty:

            return (
                None,
                "### 🟡 No matching records to export."
            )

        # ----------------------------------------------------
        # OUTPUT DIRECTORY
        # ----------------------------------------------------

        OUTPUT_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        # ----------------------------------------------------
        # FILE NAME
        # ----------------------------------------------------

        timestamp = (
            datetime.now()
            .strftime(
                "%Y%m%d_%H%M%S"
            )
        )

        export_file = (
            OUTPUT_DIR
            / f"filtered_roster_{timestamp}.xlsx"
        )

        # ----------------------------------------------------
        # PREPARE DISPLAY
        # ----------------------------------------------------

        export_df = df.copy()

        if "Date" in export_df.columns:

            export_df["Date"] = (
                export_df["Date"]
                .dt.strftime(
                    "%d/%m/%Y"
                )
            )

        # ----------------------------------------------------
        # WRITE EXCEL WITH PROFESSIONAL FORMATTING
        # ----------------------------------------------------

        with pd.ExcelWriter(
            export_file,
            engine="openpyxl"
        ) as writer:

            export_df.to_excel(
                writer,
                index=False,
                sheet_name="Filtered Roster"
            )

            worksheet = (
                writer.sheets[
                    "Filtered Roster"
                ]
            )

            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.page import PageMargins

            header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
            header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
            thin_border = Border(
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
                top=Side(style="thin", color="D9D9D9"),
                bottom=Side(style="thin", color="D9D9D9")
            )

            primary_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
            cover_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
            unassigned_fill = PatternFill(fill_type="solid", fgColor="F4CCCC")

            worksheet.sheet_view.showGridLines = False
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

            if worksheet.max_row >= 1:
                worksheet.row_dimensions[1].height = 28
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border

            assignment_col_idx = None
            for idx, col_name in enumerate(export_df.columns, start=1):
                if col_name == "Assignment Type":
                    assignment_col_idx = idx
                    break

            for row in range(2, worksheet.max_row + 1):
                worksheet.row_dimensions[row].height = 22
                
                assign_type_val = ""
                if assignment_col_idx:
                    assign_type_val = str(worksheet.cell(row=row, column=assignment_col_idx).value or "").upper()

                for col in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.font = Font(name="Calibri", size=10)

                    if assign_type_val == "PRIMARY":
                        cell.fill = primary_fill
                    elif assign_type_val == "COVER":
                        cell.fill = cover_fill
                    elif assign_type_val == "UNASSIGNED":
                        cell.fill = unassigned_fill

            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter
                for cell in column_cells:
                    try:
                        cell_length = len(str(cell.value or ""))
                        if cell_length > max_length:
                            max_length = cell_length
                    except Exception:
                        pass
                worksheet.column_dimensions[column_letter].width = min(max(max_length + 4, 15), 45)

            worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
            worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
            worksheet.page_setup.fitToWidth = 1
            worksheet.page_setup.fitToHeight = 0
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            worksheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5)

        return (
            str(export_file),

            "### 🟢 FILTERED EXCEL EXPORTED SUCCESSFULLY (PROFESSIONAL FORMAT)\n\n"
            f"Records exported: **{len(export_df)}**\n\n"
            f"File:\n`{export_file}`"
        )

    except Exception as e:

        return (
            None,
            "### 🔴 EXCEL EXPORT FAILED\n\n"
            f"```text\n{e}\n```"
        )


# ============================================================
# STEP 7.6 - STAFF & DEPARTMENT MANAGEMENT
# ============================================================

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill


def _clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _normalize_yes_no(value, default="NO"):
    value = _clean_text(value).upper()
    if value in {"YES", "NO"}:
        return value
    return default


def _read_master_dataframe(file_path):
    df, error = safe_read_excel(file_path)
    if error:
        return pd.DataFrame(), error
    return df, None


def _save_workbook_safely(file_path, workbook):
    try:
        workbook.save(file_path)
        return True, None
    except PermissionError:
        return (
            False,
            f"Permission denied: {file_path.name}. "
            "Please close the Excel file and try again."
        )
    except OSError as e:
        return False, str(e)
    except Exception as e:
        return False, str(e)


def _format_master_sheet(ws):
    if ws.max_column == 0:
        return

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )
    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            try:
                max_length = max(
                    max_length,
                    len(str(cell.value or ""))
                )
            except Exception:
                pass

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            35
        )


def _get_staff_and_duty_tables():
    staff_df, staff_error = _read_master_dataframe(
        STAFF_MASTER_FILE
    )
    duty_df, duty_error = _read_master_dataframe(
        DUTY_MASTER_FILE
    )

    if staff_error:
        staff_df = pd.DataFrame()

    if duty_error:
        duty_df = pd.DataFrame()

    return staff_df, duty_df


def get_master_tables():
    return _get_staff_and_duty_tables()


def get_master_duty_choices():
    """Return duties that actually exist in Duty Master."""
    duty_df, error = _read_master_dataframe(
        DUTY_MASTER_FILE
    )

    if error or duty_df.empty or "Duty" not in duty_df.columns:
        return []

    values = (
        duty_df["Duty"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    return sorted(
        [value for value in values.tolist() if value],
        key=str.casefold
    )


def get_employee_duty_choices():
    """
    Return all duty-permission choices for employee management.

    Existing master duties are always preserved.  The standard RKM
    duties are also shown so the user can select Morning/Afternoon/Night
    etc. directly.  If a selected standard duty does not yet exist in the
    master files, add_new_employee() creates it safely before saving the
    employee permission.
    """
    values = set(get_master_duty_choices())
    values.update(STANDARD_RKM_DUTIES)
    return sorted(
        [value for value in values if _clean_text(value)],
        key=str.casefold
    )


def _ensure_staff_workbook_structure(workbook):
    ws = workbook.active

    headers = [
        _clean_text(ws.cell(1, col).value)
        for col in range(1, ws.max_column + 1)
    ]

    while headers and headers[-1] == "":
        headers.pop()

    if not headers:
        headers = ["Staff Name", "Category", "Active"]
        for col, header in enumerate(headers, start=1):
            ws.cell(1, col).value = header
    else:
        if "Staff Name" not in headers:
            headers.append("Staff Name")
            ws.cell(1, len(headers)).value = "Staff Name"

        if "Category" not in headers:
            headers.append("Category")
            ws.cell(1, len(headers)).value = "Category"

        if "Active" not in headers:
            headers.append("Active")
            ws.cell(1, len(headers)).value = "Active"

    return ws, headers


def _ensure_duty_workbook_structure(workbook):
    ws = workbook.active

    headers = [
        _clean_text(ws.cell(1, col).value)
        for col in range(1, ws.max_column + 1)
    ]

    while headers and headers[-1] == "":
        headers.pop()

    if not headers:
        headers = ["Duty", "Required Staff"]
        for col, header in enumerate(headers, start=1):
            ws.cell(1, col).value = header
    else:
        if "Duty" not in headers:
            headers.append("Duty")
            ws.cell(1, len(headers)).value = "Duty"

        if "Required Staff" not in headers:
            headers.append("Required Staff")
            ws.cell(1, len(headers)).value = "Required Staff"

    return ws, headers


def add_new_employee(
    staff_name,
    category,
    active_value,
    selected_duties
):
    """Add an employee and safely create any newly selected standard duties."""
    name = _clean_text(staff_name)
    category = _clean_text(category)
    active_value = _normalize_yes_no(active_value, "YES")

    current_staff, current_duty = _get_staff_and_duty_tables()

    def result(message, duties=None):
        choices = get_employee_duty_choices() if duties is None else duties
        return (
            message,
            current_staff,
            current_duty,
            gr.update(
                choices=choices,
                value=[]
            )
        )

    if not name:
        return result(
            "### 🔴 EMPLOYEE NOT ADDED\n\nStaff Name is required."
        )

    if not category:
        return result(
            "### 🔴 EMPLOYEE NOT ADDED\n\nCategory is required."
        )

    selected_duties = {
        _clean_text(value)
        for value in (selected_duties or [])
        if _clean_text(value)
    }

    try:
        # ------------------------------------------------------------
        # IMPORTANT: if the user selected Morning/Afternoon/Night/etc.
        # and those duties are not yet in Duty Master, create them first.
        # This makes the checkbox options actually functional.
        # ------------------------------------------------------------
        existing_duties = set(get_master_duty_choices())
        missing_selected_duties = [
            duty for duty in sorted(selected_duties, key=str.casefold)
            if duty not in existing_duties
        ]

        if missing_selected_duties:
            for duty in missing_selected_duties:
                add_result = add_new_department(duty, 1)
                if not str(add_result[0]).startswith("### 🟢 DEPARTMENT / DUTY ADDED SUCCESSFULLY"):
                    return (
                        "### 🔴 EMPLOYEE NOT ADDED\n\n"
                        f"Could not create duty **{duty}** first.\n\n"
                        f"{add_result[0]}"
                    ,
                        *_get_staff_and_duty_tables(),
                        gr.update(choices=get_employee_duty_choices(), value=[])
                    )

        STAFF_MASTER_FILE.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        if STAFF_MASTER_FILE.exists():
            workbook = load_workbook(STAFF_MASTER_FILE)
        else:
            from openpyxl import Workbook
            workbook = Workbook()

        ws, headers = _ensure_staff_workbook_structure(workbook)

        name_col = headers.index("Staff Name") + 1
        existing_names = set()

        for row in range(2, ws.max_row + 1):
            value = _clean_text(ws.cell(row, name_col).value)
            if value:
                existing_names.add(value.casefold())

        if name.casefold() in existing_names:
            return result(
                "### 🟡 EMPLOYEE NOT ADDED\n\n"
                f"**{name}** already exists in Staff Master."
            )

        # Use the complete duty list after any auto-created duties.
        duty_choices = get_master_duty_choices()

        for duty in duty_choices:
            if duty not in headers:
                headers.append(duty)
                ws.cell(1, len(headers)).value = duty

                for row in range(2, ws.max_row + 1):
                    ws.cell(row, len(headers)).value = "NO"

        new_row = ws.max_row + 1

        for col, header in enumerate(headers, start=1):
            if header == "Staff Name":
                value = name
            elif header == "Category":
                value = category
            elif header == "Active":
                value = active_value
            elif header in duty_choices:
                value = "YES" if header in selected_duties else "NO"
            else:
                value = ""

            ws.cell(new_row, col).value = value

        _format_master_sheet(ws)

        ok, error = _save_workbook_safely(
            STAFF_MASTER_FILE,
            workbook
        )

        if not ok:
            return result(
                "### 🔴 EMPLOYEE NOT ADDED\n\n"
                f"{error}"
            )

        staff_df, duty_df = _get_staff_and_duty_tables()
        choices = get_employee_duty_choices()

        return (
            "### 🟢 EMPLOYEE ADDED SUCCESSFULLY\n\n"
            f"**{name}** has been added to Staff Master.\n\n"
            f"Active: **{active_value}**  \n"
            f"Category: **{category}**  \n"
            f"Duty permissions: **{len(selected_duties)}**",
            staff_df,
            duty_df,
            gr.update(choices=choices, value=[])
        )

    except PermissionError:
        return result(
            "### 🔴 EMPLOYEE NOT ADDED\n\n"
            "Staff Master is open in Excel. Please close it and try again."
        )
    except Exception as e:
        return result(
            "### 🔴 EMPLOYEE NOT ADDED\n\n"
            f"```text\n{e}\n```"
        )


def add_new_department(
    duty_name,
    required_staff
):
    duty_name = _clean_text(duty_name)

    current_staff, current_duty = _get_staff_and_duty_tables()

    def result(message, choices=None):
        return (
            message,
            current_staff,
            current_duty,
            gr.update(
                choices=get_master_duty_choices() if choices is None else choices,
                value=[]
            )
        )

    if not duty_name:
        return result(
            "### 🔴 DEPARTMENT / DUTY NOT ADDED\n\n"
            "Department / Duty name is required."
        )

    try:
        required_staff = int(required_staff)
    except (TypeError, ValueError):
        required_staff = 0

    if required_staff < 1:
        return result(
            "### 🔴 DEPARTMENT / DUTY NOT ADDED\n\n"
            "Required Staff must be **1 or more**."
        )

    try:
        DUTY_MASTER_FILE.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        if DUTY_MASTER_FILE.exists():
            duty_workbook = load_workbook(DUTY_MASTER_FILE)
        else:
            from openpyxl import Workbook
            duty_workbook = Workbook()

        duty_ws, duty_headers = _ensure_duty_workbook_structure(
            duty_workbook
        )

        duty_col = duty_headers.index("Duty") + 1
        existing_duties = set()

        for row in range(2, duty_ws.max_row + 1):
            value = _clean_text(duty_ws.cell(row, duty_col).value)
            if value:
                existing_duties.add(value.casefold())

        if duty_name.casefold() in existing_duties:
            return result(
                "### 🟡 DEPARTMENT / DUTY NOT ADDED\n\n"
                f"**{duty_name}** already exists in Duty Master."
            )

        new_row = duty_ws.max_row + 1
        for col, header in enumerate(duty_headers, start=1):
            if header == "Duty":
                value = duty_name
            elif header == "Required Staff":
                value = required_staff
            else:
                value = ""
            duty_ws.cell(new_row, col).value = value

        _format_master_sheet(duty_ws)

        # First make sure Staff Master can accept the new permission column.
        STAFF_MASTER_FILE.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        if STAFF_MASTER_FILE.exists():
            staff_workbook = load_workbook(STAFF_MASTER_FILE)
        else:
            from openpyxl import Workbook
            staff_workbook = Workbook()

        staff_ws, staff_headers = _ensure_staff_workbook_structure(
            staff_workbook
        )

        if duty_name not in staff_headers:
            new_permission_col = len(staff_headers) + 1
            staff_ws.cell(1, new_permission_col).value = duty_name

            for row in range(2, staff_ws.max_row + 1):
                staff_ws.cell(row, new_permission_col).value = "NO"

        _format_master_sheet(staff_ws)

        # Save Staff Master first. If it is locked, Duty Master is not changed.
        ok, error = _save_workbook_safely(
            STAFF_MASTER_FILE,
            staff_workbook
        )

        if not ok:
            return result(
                "### 🔴 DEPARTMENT / DUTY NOT ADDED\n\n"
                f"{error}\n\n"
                "Close Staff Master in Excel and try again."
            )

        ok, error = _save_workbook_safely(
            DUTY_MASTER_FILE,
            duty_workbook
        )

        if not ok:
            return result(
                "### 🟠 STAFF PERMISSION COLUMN UPDATED, BUT DUTY MASTER "
                "COULD NOT BE SAVED\n\n"
                f"{error}\n\n"
                "Close Duty Master in Excel and try again."
            )

        staff_df, duty_df = _get_staff_and_duty_tables()
        choices = get_master_duty_choices()

        return (
            "### 🟢 DEPARTMENT / DUTY ADDED SUCCESSFULLY\n\n"
            f"**{duty_name}** has been added.\n\n"
            f"Required Staff: **{required_staff}**\n\n"
            "Existing employees received **NO** permission for this new duty."
            " You can grant permission by editing Staff Master or by using "
            "the employee-add workflow for future employees.",
            staff_df,
            duty_df,
            gr.update(choices=choices, value=[])
        )

    except PermissionError:
        return result(
            "### 🔴 DEPARTMENT / DUTY NOT ADDED\n\n"
            "One of the master Excel files is open. Please close "
            "**staff_master.xlsx** and **duty_master.xlsx**, then try again."
        )
    except Exception as e:
        return result(
            "### 🔴 DEPARTMENT / DUTY NOT ADDED\n\n"
            f"```text\n{e}\n```"
        )


def upload_duty_master_excel(uploaded_file):
    """Validate and safely replace Duty Master from an uploaded Excel file.

    The existing Duty Master is backed up before replacement.  The Staff Master
    is preserved and automatically receives missing permission columns as NO.
    """
    staff_df, duty_df = _get_staff_and_duty_tables()

    if not uploaded_file:
        return (
            "### 🟡 NO DUTY EXCEL SELECTED\n\nSelect a `.xlsx` or `.xls` Duty Master file first.",
            staff_df,
            duty_df,
            gr.update(choices=get_employee_duty_choices(), value=[]),
        )

    source = Path(str(uploaded_file))
    if not source.exists():
        return (
            "### 🔴 DUTY MASTER UPLOAD FAILED\n\nUploaded file was not found.",
            staff_df,
            duty_df,
            gr.update(choices=get_employee_duty_choices(), value=[]),
        )

    try:
        uploaded_df, read_error = safe_read_excel(source)
        if read_error:
            raise ValueError(f"Could not read uploaded Excel: {read_error}")

        required = ["Duty", "Required Staff"]
        missing = [c for c in required if c not in uploaded_df.columns]
        if missing:
            raise ValueError(
                "Missing required column(s): " + ", ".join(missing)
            )

        check = uploaded_df[["Duty", "Required Staff"]].copy()
        check["Duty"] = check["Duty"].fillna("").astype(str).str.strip()
        if (check["Duty"] == "").any():
            raise ValueError("Duty column contains blank duty names.")

        normalized_names = check["Duty"].str.casefold()
        duplicates = check.loc[normalized_names.duplicated(keep=False), "Duty"].tolist()
        if duplicates:
            unique_dupes = sorted(set(duplicates), key=str.casefold)
            raise ValueError(
                "Duplicate duty name(s) found: " + ", ".join(unique_dupes)
            )

        numeric_required = pd.to_numeric(
            check["Required Staff"], errors="coerce"
        )
        if numeric_required.isna().any():
            raise ValueError("Required Staff must contain valid numbers.")
        if (numeric_required < 1).any():
            raise ValueError("Required Staff must be 1 or more for every duty.")

        # Preserve the uploaded workbook's additional columns, but ensure the
        # two engine-required columns are present and clean.
        uploaded_df["Duty"] = uploaded_df["Duty"].fillna("").astype(str).str.strip()
        uploaded_df["Required Staff"] = numeric_required.astype(int)

        DUTY_MASTER_FILE.parent.mkdir(parents=True, exist_ok=True)
        STAFF_MASTER_FILE.parent.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        duty_backup = None
        staff_backup = None

        if DUTY_MASTER_FILE.exists():
            duty_backup = DUTY_MASTER_FILE.with_name(
                f"duty_master_backup_{timestamp}.xlsx"
            )
            shutil.copy2(DUTY_MASTER_FILE, duty_backup)

        if STAFF_MASTER_FILE.exists():
            staff_backup = STAFF_MASTER_FILE.with_name(
                f"staff_master_backup_{timestamp}.xlsx"
            )
            shutil.copy2(STAFF_MASTER_FILE, staff_backup)

        # Update Staff Master permission columns BEFORE replacing Duty Master.
        # If this fails, the original Duty Master remains untouched.
        if STAFF_MASTER_FILE.exists():
            staff_workbook = load_workbook(STAFF_MASTER_FILE)
        else:
            from openpyxl import Workbook
            staff_workbook = Workbook()

        staff_ws, staff_headers = _ensure_staff_workbook_structure(staff_workbook)

        uploaded_duties = [
            _clean_text(v) for v in uploaded_df["Duty"].tolist()
            if _clean_text(v)
        ]

        for duty in uploaded_duties:
            existing_header = None
            for header in staff_headers:
                if _clean_text(header).casefold() == duty.casefold():
                    existing_header = header
                    break

            if existing_header is None:
                new_col = len(staff_headers) + 1
                staff_ws.cell(1, new_col).value = duty
                staff_headers.append(duty)
                for row in range(2, staff_ws.max_row + 1):
                    staff_ws.cell(row, new_col).value = "NO"

        _format_master_sheet(staff_ws)
        ok, error = _save_workbook_safely(STAFF_MASTER_FILE, staff_workbook)
        if not ok:
            raise PermissionError(error)

        # Write the uploaded Duty Master through pandas so the user's uploaded
        # sheet data is retained instead of rebuilding it from scratch.
        uploaded_df.to_excel(DUTY_MASTER_FILE, index=False)

        # Re-format the Duty Master with openpyxl after pandas writes it.
        duty_workbook = load_workbook(DUTY_MASTER_FILE)
        _format_master_sheet(duty_workbook.active)
        ok, error = _save_workbook_safely(DUTY_MASTER_FILE, duty_workbook)
        if not ok:
            raise PermissionError(error)

        new_staff_df, new_duty_df = _get_staff_and_duty_tables()
        backup_text = ""
        if duty_backup:
            backup_text += f"\nDuty backup: `{duty_backup.name}`"
        if staff_backup:
            backup_text += f"\nStaff backup: `{staff_backup.name}`"

        return (
            "### 🟢 DUTY MASTER UPLOADED SUCCESSFULLY\n\n"
            f"Imported **{len(new_duty_df)} duty/duties** from **{source.name}**."
            "\n\nExisting staff data was preserved. Missing duty-permission columns were"
            " added as **NO**."
            f"\n\n{backup_text.strip()}",
            new_staff_df,
            new_duty_df,
            gr.update(choices=get_employee_duty_choices(), value=[]),
        )

    except PermissionError as e:
        return (
            "### 🔴 DUTY MASTER UPLOAD FAILED\n\n"
            f"{e}\n\nPlease close the master Excel files and try again.",
            *_get_staff_and_duty_tables(),
            gr.update(choices=get_employee_duty_choices(), value=[]),
        )
    except Exception as e:
        # Best-effort rollback if a new Duty Master was written but failed
        # during final formatting.
        try:
            if 'duty_backup' in locals() and duty_backup and duty_backup.exists():
                shutil.copy2(duty_backup, DUTY_MASTER_FILE)
            if 'staff_backup' in locals() and staff_backup and staff_backup.exists():
                shutil.copy2(staff_backup, STAFF_MASTER_FILE)
        except Exception:
            pass

        return (
            "### 🔴 DUTY MASTER UPLOAD FAILED\n\n"
            f"```text\n{e}\n```",
            *_get_staff_and_duty_tables(),
            gr.update(choices=get_employee_duty_choices(), value=[]),
        )


def add_standard_rkm_duties(selected_duties):
    """Add selected standard RKM duties to Duty Master."""
    selected = []
    for value in (selected_duties or []):
        value = _clean_text(value)
        if value and value in STANDARD_RKM_DUTIES and value not in selected:
            selected.append(value)

    staff_df, duty_df = _get_staff_and_duty_tables()

    if not selected:
        return (
            "### 🟡 NO STANDARD DUTY SELECTED\n\nSelect one or more standard duties first.",
            staff_df,
            duty_df,
            gr.update(choices=get_master_duty_choices(), value=[]),
        )

    messages = []
    for duty in selected:
        result = add_new_department(duty, 1)
        messages.append(result[0])

    staff_df, duty_df = _get_staff_and_duty_tables()
    choices = get_master_duty_choices()

    return (
        "### 🟢 STANDARD DUTY UPDATE COMPLETE\n\n" + "\n\n".join(messages),
        staff_df,
        duty_df,
        gr.update(choices=choices, value=[]),
    )


def refresh_master_data():
    staff_df, duty_df = _get_staff_and_duty_tables()
    choices = get_master_duty_choices()

    return (
        staff_df,
        duty_df,
        gr.update(choices=choices, value=[]),
        "### 🟢 MASTER DATA REFRESHED SUCCESSFULLY"
    )



# ============================================================
# STEP 7.6 - EMPLOYEE / DEPARTMENT VIEW EDIT DELETE
# ============================================================

def _staff_names_from_df(staff_df):
    if staff_df is None or staff_df.empty or "Staff Name" not in staff_df.columns:
        return []
    return sorted(
        [
            _clean_text(v)
            for v in staff_df["Staff Name"].tolist()
            if _clean_text(v)
        ],
        key=str.casefold,
    )


def _duty_names_from_df(duty_df):
    if duty_df is None or duty_df.empty or "Duty" not in duty_df.columns:
        return []
    return sorted(
        [
            _clean_text(v)
            for v in duty_df["Duty"].tolist()
            if _clean_text(v)
        ],
        key=str.casefold,
    )


def _staff_row_to_details(staff_df, staff_name):
    if staff_df is None or staff_df.empty or "Staff Name" not in staff_df.columns:
        return None

    target = _clean_text(staff_name).casefold()
    for _, row in staff_df.iterrows():
        if _clean_text(row.get("Staff Name", "")).casefold() == target:
            return row
    return None


def view_employee_details(staff_name):
    staff_df, duty_df = _get_staff_and_duty_tables()
    names = _staff_names_from_df(staff_df)

    if not _clean_text(staff_name):
        return (
            "### ⚪ Select an employee and click **VIEW**.",
            gr.update(choices=names, value=None),
            "",
            "",
            "YES",
            gr.update(choices=_duty_names_from_df(duty_df), value=[]),
            False,
        )

    row = _staff_row_to_details(staff_df, staff_name)
    if row is None:
        return (
            "### 🔴 EMPLOYEE NOT FOUND\n\nPlease refresh Master Data and try again.",
            gr.update(choices=names, value=None),
            "",
            "",
            "YES",
            gr.update(choices=_duty_names_from_df(duty_df), value=[]),
            False,
        )

    duties = get_employee_duty_choices()
    permitted = []
    for duty in duties:
        if duty in staff_df.columns and _normalize_yes_no(row.get(duty), "NO") == "YES":
            permitted.append(duty)

    name = _clean_text(row.get("Staff Name", ""))
    category = _clean_text(row.get("Category", ""))
    active = _normalize_yes_no(row.get("Active", "YES"), "YES")

    view_text = (
        "### 🟢 EMPLOYEE DETAILS\n\n"
        f"**Name:** {name}  \n"
        f"**Category:** {category or '—'}  \n"
        f"**Active:** {active}  \n"
        f"**Duty Permissions:** {', '.join(permitted) if permitted else 'None'}\n\n"
        "You can now edit the fields below."
    )

    return (
        view_text,
        gr.update(choices=names, value=name),
        name,
        category,
        active,
        gr.update(choices=duties, value=permitted),
        False,
    )


def edit_employee_details(original_name, new_name, category, active_value, selected_duties):
    staff_df, duty_df = _get_staff_and_duty_tables()
    names = _staff_names_from_df(staff_df)
    duties = get_employee_duty_choices()

    original_name = _clean_text(original_name)
    new_name = _clean_text(new_name)
    category = _clean_text(category)
    active_value = _normalize_yes_no(active_value, "YES")
    selected = {_clean_text(v) for v in (selected_duties or []) if _clean_text(v)}

    if not original_name:
        return (
            "### 🔴 EMPLOYEE NOT SAVED\n\nSelect an employee and click **VIEW** first.",
            staff_df,
            duty_df,
            gr.update(choices=names, value=None),
            gr.update(choices=duties, value=[]),
            False,
            "",
            "",
            "YES",
            gr.update(choices=duties, value=[]),
        )

    if not new_name:
        return (
            "### 🔴 EMPLOYEE NOT SAVED\n\nEmployee Name is required.",
            staff_df,
            duty_df,
            gr.update(choices=names, value=original_name),
            gr.update(choices=duties, value=[]),
            False,
            new_name,
            category,
            active_value,
            gr.update(choices=duties, value=list(selected)),
        )

    if not category:
        return (
            "### 🔴 EMPLOYEE NOT SAVED\n\nCategory is required.",
            staff_df,
            duty_df,
            gr.update(choices=names, value=original_name),
            gr.update(choices=duties, value=[]),
            False,
            new_name,
            category,
            active_value,
            gr.update(choices=duties, value=list(selected)),
        )

    try:
        wb = load_workbook(STAFF_MASTER_FILE)
        ws, headers = _ensure_staff_workbook_structure(wb)
        name_col = headers.index("Staff Name") + 1
        category_col = headers.index("Category") + 1
        active_col = headers.index("Active") + 1

        target_row = None
        for r in range(2, ws.max_row + 1):
            if _clean_text(ws.cell(r, name_col).value).casefold() == original_name.casefold():
                target_row = r
                break

        if target_row is None:
            return (
                "### 🔴 EMPLOYEE NOT FOUND\n\nRefresh Master Data and try again.",
                staff_df,
                duty_df,
                gr.update(choices=names, value=None),
                gr.update(choices=duties, value=[]),
                False,
                "",
                "",
                "YES",
                gr.update(choices=duties, value=[]),
            )

        for r in range(2, ws.max_row + 1):
            if r == target_row:
                continue
            existing = _clean_text(ws.cell(r, name_col).value)
            if existing and existing.casefold() == new_name.casefold():
                return (
                    f"### 🟡 EMPLOYEE NOT SAVED\n\n**{new_name}** already exists.",
                    staff_df,
                    duty_df,
                    gr.update(choices=names, value=original_name),
                    gr.update(choices=duties, value=[]),
                    False,
                    new_name,
                    category,
                    active_value,
                    gr.update(choices=duties, value=list(selected)),
                )

        ws.cell(target_row, name_col).value = new_name
        ws.cell(target_row, category_col).value = category
        ws.cell(target_row, active_col).value = active_value

        # Update every known duty permission while preserving any extra columns.
        # Standard RKM duties are allowed here too.
        for duty in duties:
            if duty not in headers:
                ws.cell(1, ws.max_column + 1).value = duty
                headers.append(duty)
            col = headers.index(duty) + 1
            ws.cell(target_row, col).value = "YES" if duty in selected else "NO"

        _format_master_sheet(ws)
        ok, error = _save_workbook_safely(STAFF_MASTER_FILE, wb)
        if not ok:
            return (
                f"### 🔴 EMPLOYEE NOT SAVED\n\n{error}",
                staff_df,
                duty_df,
                gr.update(choices=names, value=original_name),
                gr.update(choices=duties, value=[]),
                False,
                new_name,
                category,
                active_value,
                gr.update(choices=duties, value=list(selected)),
            )

        new_staff_df, new_duty_df = _get_staff_and_duty_tables()
        new_names = _staff_names_from_df(new_staff_df)
        new_duties = _duty_names_from_df(new_duty_df)
        row = _staff_row_to_details(new_staff_df, new_name)
        permitted = [
            duty for duty in new_duties
            if row is not None and duty in new_staff_df.columns
            and _normalize_yes_no(row.get(duty), "NO") == "YES"
        ]

        return (
            f"### 🟢 EMPLOYEE UPDATED SUCCESSFULLY\n\n**{new_name}** has been updated.",
            new_staff_df,
            new_duty_df,
            gr.update(choices=new_names, value=new_name),
            gr.update(choices=new_duties, value=[]),
            False,
            new_name,
            category,
            active_value,
            gr.update(choices=new_duties, value=permitted),
        )

    except PermissionError:
        return (
            "### 🔴 EMPLOYEE NOT SAVED\n\nClose **staff_master.xlsx** in Excel and try again.",
            staff_df,
            duty_df,
            gr.update(choices=names, value=original_name),
            gr.update(choices=duties, value=[]),
            False,
            new_name,
            category,
            active_value,
            gr.update(choices=duties, value=list(selected)),
        )
    except Exception as e:
        return (
            f"### 🔴 EMPLOYEE EDIT FAILED\n\n```text\n{e}\n```",
            staff_df,
            duty_df,
            gr.update(choices=names, value=original_name),
            gr.update(choices=duties, value=[]),
            False,
            new_name,
            category,
            active_value,
            gr.update(choices=duties, value=list(selected)),
        )


def delete_employee_details(staff_name, confirmed):
    staff_df, duty_df = _get_staff_and_duty_tables()
    names = _staff_names_from_df(staff_df)
    duties = _duty_names_from_df(duty_df)
    name = _clean_text(staff_name)

    base_return = lambda msg, n=names, d=duties: (
        msg,
        staff_df,
        duty_df,
        gr.update(choices=n, value=None),
        gr.update(choices=d, value=[]),
        "",
        "",
        "",
        "YES",
        gr.update(choices=d, value=[]),
        False,
    )

    if not name:
        return base_return("### 🔴 EMPLOYEE NOT DELETED\n\nSelect an employee first.")

    if not confirmed:
        return base_return(
            "### 🟠 DELETE BLOCKED\n\n"
            "You must first click **VIEW** and tick **I have viewed the details and confirm DELETE**."
        )

    try:
        wb = load_workbook(STAFF_MASTER_FILE)
        ws, headers = _ensure_staff_workbook_structure(wb)
        name_col = headers.index("Staff Name") + 1
        target_row = None
        for r in range(2, ws.max_row + 1):
            if _clean_text(ws.cell(r, name_col).value).casefold() == name.casefold():
                target_row = r
                break

        if target_row is None:
            return base_return("### 🔴 EMPLOYEE NOT FOUND\n\nRefresh Master Data and try again.")

        ws.delete_rows(target_row, 1)
        _format_master_sheet(ws)
        ok, error = _save_workbook_safely(STAFF_MASTER_FILE, wb)
        if not ok:
            return base_return(f"### 🔴 EMPLOYEE NOT DELETED\n\n{error}")

        new_staff_df, new_duty_df = _get_staff_and_duty_tables()
        new_names = _staff_names_from_df(new_staff_df)
        new_duties = _duty_names_from_df(new_duty_df)
        return (
            f"### 🟢 EMPLOYEE DELETED SUCCESSFULLY\n\n**{name}** has been removed from Staff Master.",
            new_staff_df,
            new_duty_df,
            gr.update(choices=new_names, value=None),
            gr.update(choices=new_duties, value=[]),
            "",
            "",
            "",
            "YES",
            gr.update(choices=new_duties, value=[]),
            False,
        )
    except PermissionError:
        return base_return("### 🔴 EMPLOYEE NOT DELETED\n\nClose **staff_master.xlsx** in Excel and try again.")
    except Exception as e:
        return base_return(f"### 🔴 EMPLOYEE DELETE FAILED\n\n```text\n{e}\n```")


def view_department_details(duty_name):
    staff_df, duty_df = _get_staff_and_duty_tables()
    duties = _duty_names_from_df(duty_df)

    if not _clean_text(duty_name):
        return (
            "### ⚪ Select a department/duty and click **VIEW**.",
            gr.update(choices=duties, value=None),
            "",
            1,
        )

    target = _clean_text(duty_name).casefold()
    row = None
    for _, r in duty_df.iterrows():
        if _clean_text(r.get("Duty", "")).casefold() == target:
            row = r
            break

    if row is None:
        return (
            "### 🔴 DEPARTMENT / DUTY NOT FOUND",
            gr.update(choices=duties, value=None),
            "",
            1,
        )

    required = row.get("Required Staff", 1)
    try:
        required = int(float(required))
    except Exception:
        required = 1

    name = _clean_text(row.get("Duty", ""))
    return (
        "### 🟢 DEPARTMENT / DUTY DETAILS\n\n"
        f"**Name:** {name}  \n"
        f"**Required Staff:** {required}  \n\n"
        "You can now edit the fields below.",
        gr.update(choices=duties, value=name),
        name,
        required,
    )


def edit_department_details(original_name, new_name, required_staff):
    """Safely edit a Duty Master row and keep Staff Master permission column in sync."""
    staff_df, duty_df = _get_staff_and_duty_tables()
    duties = _duty_names_from_df(duty_df)

    original_name = _clean_text(original_name)
    new_name = _clean_text(new_name)

    try:
        required_staff = int(float(required_staff))
    except (TypeError, ValueError):
        required_staff = 0

    def fail(message, keep_name=None, keep_required=None):
        keep_name = original_name if keep_name is None else keep_name
        keep_required = max(required_staff, 1) if keep_required is None else keep_required
        return (
            message,
            staff_df,
            duty_df,
            gr.update(choices=duties, value=original_name or None),
            gr.update(choices=duties, value=[]),
            original_name,
            keep_name,
            keep_required,
        )

    if not original_name:
        return fail(
            "### 🔴 DEPARTMENT / DUTY NOT SAVED\n\n"
            "Select a department/duty and click **VIEW** first.",
            keep_name="",
            keep_required=1,
        )

    if not new_name:
        return fail(
            "### 🔴 DEPARTMENT / DUTY NOT SAVED\n\n"
            "Department / Duty name is required."
        )

    if required_staff < 1:
        return fail(
            "### 🔴 DEPARTMENT / DUTY NOT SAVED\n\n"
            "Required Staff must be **1 or more**.",
            keep_name=new_name,
            keep_required=1,
        )

    # Work only with real Excel files.
    if not DUTY_MASTER_FILE.exists():
        return fail(f"### 🔴 DEPARTMENT / DUTY NOT SAVED\n\nMissing file: `{DUTY_MASTER_FILE}`")
    if not STAFF_MASTER_FILE.exists():
        return fail(f"### 🔴 DEPARTMENT / DUTY NOT SAVED\n\nMissing file: `{STAFF_MASTER_FILE}`")

    try:
        # ========================================================
        # 1. READ DUTY MASTER
        # ========================================================
        wb = load_workbook(DUTY_MASTER_FILE)
        ws, headers = _ensure_duty_workbook_structure(wb)

        # Case-insensitive header lookup prevents edit failures after
        # Excel files have been created/edited by another program.
        duty_header_map = {
            _clean_text(h).casefold(): idx + 1
            for idx, h in enumerate(headers)
            if _clean_text(h)
        }
        duty_col = duty_header_map.get("duty")
        required_col = duty_header_map.get("required staff")

        if not duty_col or not required_col:
            return fail(
                "### 🔴 DEPARTMENT / DUTY NOT SAVED\n\n"
                "Duty Master must contain **Duty** and **Required Staff** columns."
            )

        target_row = None
        for r in range(2, ws.max_row + 1):
            cell_value = _clean_text(ws.cell(r, duty_col).value)
            if cell_value and cell_value.casefold() == original_name.casefold():
                target_row = r
                break

        if target_row is None:
            return fail(
                "### 🔴 DEPARTMENT / DUTY NOT FOUND\n\n"
                f"Could not find **{original_name}** in Duty Master.\n\n"
                "Click **REFRESH MASTER DATA**, select the duty again, and click **VIEW**."
            )

        # Do not allow duplicate duty names.
        for r in range(2, ws.max_row + 1):
            if r == target_row:
                continue
            existing = _clean_text(ws.cell(r, duty_col).value)
            if existing and existing.casefold() == new_name.casefold():
                return fail(
                    f"### 🟡 DEPARTMENT / DUTY NOT SAVED\n\n"
                    f"**{new_name}** already exists.",
                    keep_name=new_name,
                    keep_required=required_staff,
                )

        # ========================================================
        # 2. READ STAFF MASTER AND FIND PERMISSION COLUMN
        # ========================================================
        swb = load_workbook(STAFF_MASTER_FILE)
        sws, sheaders = _ensure_staff_workbook_structure(swb)

        staff_header_map = {
            _clean_text(h).casefold(): idx + 1
            for idx, h in enumerate(sheaders)
            if _clean_text(h)
        }

        old_permission_col = staff_header_map.get(original_name.casefold())
        new_permission_col = staff_header_map.get(new_name.casefold())

        # If the duty name changes, rename the existing permission column.
        # If the old column is missing, create the new one with NO for all staff.
        if old_permission_col:
            if new_permission_col and new_permission_col != old_permission_col:
                return fail(
                    f"### 🟡 DEPARTMENT / DUTY NOT SAVED\n\n"
                    f"Staff Master already contains a permission column named **{new_name}**.",
                    keep_name=new_name,
                    keep_required=required_staff,
                )
            sws.cell(1, old_permission_col).value = new_name
        elif not new_permission_col:
            new_permission_col = sws.max_column + 1
            sws.cell(1, new_permission_col).value = new_name
            for r in range(2, sws.max_row + 1):
                sws.cell(r, new_permission_col).value = "NO"

        # ========================================================
        # 3. UPDATE BOTH WORKBOOKS IN MEMORY
        # ========================================================
        ws.cell(target_row, duty_col).value = new_name
        ws.cell(target_row, required_col).value = required_staff

        _format_master_sheet(ws)
        _format_master_sheet(sws)

        # ========================================================
        # 4. BACK UP BOTH FILES BEFORE WRITING
        # ========================================================
        staff_backup = STAFF_MASTER_FILE.with_name(
            STAFF_MASTER_FILE.stem + "_edit_backup.xlsx"
        )
        duty_backup = DUTY_MASTER_FILE.with_name(
            DUTY_MASTER_FILE.stem + "_edit_backup.xlsx"
        )

        try:
            shutil.copy2(STAFF_MASTER_FILE, staff_backup)
            shutil.copy2(DUTY_MASTER_FILE, duty_backup)
        except Exception:
            # Backups are helpful but must not prevent a valid edit.
            pass

        # ========================================================
        # 5. SAVE DUTY MASTER FIRST
        # ========================================================
        ok, error = _save_workbook_safely(DUTY_MASTER_FILE, wb)
        if not ok:
            return fail(
                "### 🔴 DEPARTMENT / DUTY NOT SAVED\n\n"
                f"Could not save **duty_master.xlsx**.\n\n{error}",
                keep_name=new_name,
                keep_required=required_staff,
            )

        # ========================================================
        # 6. SAVE STAFF MASTER
        # ========================================================
        ok, error = _save_workbook_safely(STAFF_MASTER_FILE, swb)
        if not ok:
            # Roll Duty Master back if Staff Master could not be saved.
            try:
                if duty_backup.exists():
                    shutil.copy2(duty_backup, DUTY_MASTER_FILE)
            except Exception:
                pass

            return fail(
                "### 🔴 DEPARTMENT / DUTY NOT SAVED\n\n"
                f"Staff Master could not be saved.\n\n{error}\n\n"
                "No partial change should remain in Duty Master.",
                keep_name=new_name,
                keep_required=required_staff,
            )

        # ========================================================
        # 7. RELOAD EVERYTHING FROM DISK
        # ========================================================
        new_staff_df, new_duty_df = _get_staff_and_duty_tables()
        new_duties = _duty_names_from_df(new_duty_df)

        return (
            "### 🟢 DEPARTMENT / DUTY UPDATED SUCCESSFULLY\n\n"
            f"**{original_name}** → **{new_name}**\n\n"
            f"Required Staff: **{required_staff}**",
            new_staff_df,
            new_duty_df,
            gr.update(choices=new_duties, value=new_name),
            gr.update(choices=new_duties, value=[]),
            new_name,
            new_name,
            required_staff,
        )

    except PermissionError:
        return fail(
            "### 🔴 DEPARTMENT / DUTY NOT SAVED\n\n"
            "Close **staff_master.xlsx** and **duty_master.xlsx** in Excel and try again.",
            keep_name=new_name,
            keep_required=required_staff,
        )
    except Exception as e:
        return fail(
            "### 🔴 DEPARTMENT / DUTY EDIT FAILED\n\n"
            f"```text\n{e}\n```\n\n"
            "If the problem continues, click **REFRESH MASTER DATA** and try again.",
            keep_name=new_name,
            keep_required=required_staff,
        )


def delete_department_details(duty_name, confirmed):
    staff_df, duty_df = _get_staff_and_duty_tables()
    duties = _duty_names_from_df(duty_df)
    name = _clean_text(duty_name)

    def fail(msg):
        return (
            msg,
            staff_df,
            duty_df,
            gr.update(choices=duties, value=None),
            gr.update(choices=duties, value=[]),
            "",
            "",
            1,
            False,
        )

    if not name:
        return fail("### 🔴 DEPARTMENT / DUTY NOT DELETED\n\nSelect a department/duty first.")
    if not confirmed:
        return fail(
            "### 🟠 DELETE BLOCKED\n\n"
            "You must first click **VIEW** and tick **I have viewed the details and confirm DELETE**."
        )

    try:
        wb = load_workbook(DUTY_MASTER_FILE)
        ws, headers = _ensure_duty_workbook_structure(wb)
        duty_col = headers.index("Duty") + 1
        target_row = None
        for r in range(2, ws.max_row + 1):
            if _clean_text(ws.cell(r, duty_col).value).casefold() == name.casefold():
                target_row = r
                break
        if target_row is None:
            return fail("### 🔴 DEPARTMENT / DUTY NOT FOUND")

        ws.delete_rows(target_row, 1)
        _format_master_sheet(ws)

        # Remove the matching permission column from Staff Master.
        swb = load_workbook(STAFF_MASTER_FILE)
        sws, sheaders = _ensure_staff_workbook_structure(swb)
        if name in sheaders:
            scol = sheaders.index(name) + 1
            sws.delete_cols(scol, 1)
        else:
            # Case-insensitive fallback.
            for idx, header in enumerate(sheaders, start=1):
                if _clean_text(header).casefold() == name.casefold():
                    sws.delete_cols(idx, 1)
                    break
        _format_master_sheet(sws)

        ok, error = _save_workbook_safely(STAFF_MASTER_FILE, swb)
        if not ok:
            return fail(f"### 🔴 DEPARTMENT / DUTY NOT DELETED\n\n{error}")
        ok, error = _save_workbook_safely(DUTY_MASTER_FILE, wb)
        if not ok:
            return (
                f"### 🟠 STAFF MASTER UPDATED, BUT DUTY MASTER COULD NOT BE SAVED\n\n{error}",
                *_get_staff_and_duty_tables(),
                gr.update(choices=duties, value=None),
                gr.update(choices=duties, value=[]),
                "",
                "",
                1,
                False,
            )

        new_staff_df, new_duty_df = _get_staff_and_duty_tables()
        new_duties = _duty_names_from_df(new_duty_df)
        return (
            f"### 🟢 DEPARTMENT / DUTY DELETED SUCCESSFULLY\n\n**{name}** has been removed.",
            new_staff_df,
            new_duty_df,
            gr.update(choices=new_duties, value=None),
            gr.update(choices=new_duties, value=[]),
            "",
            "",
            1,
            False,
        )
    except PermissionError:
        return fail("### 🔴 DEPARTMENT / DUTY NOT DELETED\n\nClose **staff_master.xlsx** and **duty_master.xlsx** in Excel and try again.")
    except Exception as e:
        return fail(f"### 🔴 DEPARTMENT / DUTY DELETE FAILED\n\n```text\n{e}\n```")


# ============================================================
# RULES
# ============================================================

RULES_MARKDOWN = """
## ⚙️ Active Roster Rules

| Rule | Status |
|---|---|
| 👥 Active staff automatically available | 🟢 ENABLED |
| 🏖️ Leave protection | 🟢 ENABLED |
| 🔐 Duty permission checking | 🟢 ENABLED |
| ⚖️ Fair duty distribution | 🟢 ENABLED |
| 1️⃣ Normal ONE duty per staff/day | 🟢 ENABLED |
| 2️⃣ Staff shortage SECOND duty | 🟢 ENABLED |
| 🟡 Second duty marked as COVER | 🟢 ENABLED |
| 🚫 Same duty on consecutive days | 🟢 ENABLED |
| 🚫 Same duty twice on same day | 🟢 ENABLED |
| ⚖️ Balanced cover distribution | 🟢 ENABLED |
| 🌙 Maximum 3 consecutive NIGHT duties | 🟢 ENABLED |
| 📊 Excel validation report | 🟢 ENABLED |
| 📗 Professional Excel formatting | 🟢 ENABLED |
"""


# ============================================================
# CUSTOM CSS
# ============================================================

CUSTOM_CSS = """

/* ============================================================
   GLOBAL LARGE / EASY-TO-READ FONT SETTINGS
   ============================================================ */

html,
body,
gradio-app {
    width: 100% !important;
    min-width: 100% !important;
    margin: 0 !important;
    padding: 0 !important;
    background: #f4f7fb;
    font-size: 18px !important;
}

.gradio-container {
    width: 96vw !important;
    max-width: 1600px !important;
    min-width: 0 !important;
    margin: 0 auto !important;
    padding-left: 20px !important;
    padding-right: 20px !important;
    box-sizing: border-box !important;
}

#login-screen {
    position: fixed !important;
    inset: 0 !important;
    z-index: 1000 !important;
    min-height: 100vh !important;
    overflow-y: auto !important;
    padding: 24px !important;
    box-sizing: border-box !important;
    background: linear-gradient(135deg, #17365d, #1f4e78 55%, #3b82a0) !important;
}

#login-screen,
#login-screen > .styler {
    height: auto !important;
    min-height: 100vh !important;
    overflow: visible !important;
}

#login-card {
    position: relative !important;
    width: min(520px, calc(100vw - 48px)) !important;
    margin: 4vh auto 24px !important;
    padding: 36px !important;
    border: 1px solid rgba(255, 255, 255, 0.35) !important;
    border-radius: 18px !important;
    background: rgba(255, 255, 255, 0.97) !important;
    box-shadow: 0 20px 55px rgba(0, 0, 0, 0.25) !important;
}

#login-card,
#login-card .styler {
    height: auto !important;
    min-height: 0 !important;
    overflow: visible !important;
}

#login-screen > .styler > #login-card {
    width: 100% !important;
    margin: 0 !important;
    padding: 0 !important;
    border: 0 !important;
    background: transparent !important;
    box-shadow: none !important;
}

#login-card #login-card {
    position: relative !important;
    width: min(520px, calc(100vw - 48px)) !important;
    margin: 4vh auto 24px !important;
    padding: 36px !important;
    box-sizing: border-box !important;
}

#login-card * {
    box-sizing: border-box !important;
    max-width: 100% !important;
}

#login-card .form,
#login-card .login-input,
#login-card .login-input label,
#login-card .login-input .input-container {
    width: 100% !important;
    max-width: 100% !important;
    margin-left: 0 !important;
    margin-right: 0 !important;
    overflow: visible !important;
}

#login-card .form {
    display: block !important;
}

#login-card .block:has(h1) {
    height: auto !important;
    min-height: 54px !important;
    overflow: visible !important;
}

#login-card .login-brand {
    display: block !important;
    width: 100% !important;
    margin: 0 0 30px !important;
    color: #172b4d !important;
    text-align: center !important;
}

#login-card .login-brand h1 {
    margin: 0 0 8px !important;
    color: #172b4d !important;
    font-size: 38px !important;
    line-height: 1.2 !important;
    white-space: nowrap !important;
}

#login-card .login-brand p {
    margin: 0 !important;
    color: #667085 !important;
    font-size: 22px !important;
    font-weight: 600 !important;
}

#login-card .login-input input {
    min-height: 64px !important;
    padding: 14px 18px !important;
    color: #172b4d !important;
    background: #ffffff !important;
    font-size: 25px !important;
    font-weight: 600 !important;
}

#login-card .login-input label {
    color: #172b4d !important;
    font-size: 21px !important;
    font-weight: 800 !important;
}

#login-card .login-input input::placeholder {
    color: #667085 !important;
    font-size: 21px !important;
    opacity: 1 !important;
}

#login-card h1 {
    font-size: 38px !important;
    line-height: 1.2 !important;
}

#login-card h2 {
    font-size: 28px !important;
    line-height: 1.25 !important;
}

#login-card .primary-button {
    min-height: 64px !important;
    margin-top: 18px !important;
    font-size: 25px !important;
}

#login-card h1,
#login-card h2,
#login-card p {
    text-align: center !important;
}

#login-error {
    min-height: 28px !important;
    text-align: center !important;
}

#logout-button {
    position: fixed !important;
    top: 14px !important;
    right: 14px !important;
    left: auto !important;
    z-index: 900 !important;
    min-width: 92px !important;
    width: 92px !important;
    max-width: 92px !important;
    min-height: 36px !important;
    padding: 6px 12px !important;
    border-radius: 8px !important;
    background: #b42318 !important;
    color: white !important;
    font-size: 14px !important;
    box-shadow: 0 5px 16px rgba(0, 0, 0, 0.18) !important;
}

#logout-button button {
    width: 92px !important;
    min-width: 92px !important;
    max-width: 92px !important;
}

/* ============================================================
   ALL NORMAL TEXT
   ============================================================ */

.gradio-container,
.gradio-container p,
.gradio-container span,
.gradio-container label,
.gradio-container li,
.gradio-container td,
.gradio-container th,
.gradio-container input,
.gradio-container textarea,
.gradio-container select,
.gradio-container button {
    font-size: 18px !important;
}

/* ============================================================
   MARKDOWN / HEADINGS
   ============================================================ */

.gradio-container h1 {
    font-size: 38px !important;
    line-height: 1.2 !important;
    font-weight: 800 !important;
}

.gradio-container h2 {
    font-size: 30px !important;
    line-height: 1.25 !important;
    font-weight: 800 !important;
}

.gradio-container h3 {
    font-size: 25px !important;
    line-height: 1.3 !important;
    font-weight: 750 !important;
}

.gradio-container h4,
.gradio-container h5,
.gradio-container h6 {
    font-size: 21px !important;
    font-weight: 700 !important;
}

.gradio-container p,
.gradio-container li {
    font-size: 18px !important;
    line-height: 1.55 !important;
}

.gradio-container strong,
.gradio-container b {
    font-weight: 750 !important;
}

/* ============================================================
   HERO
   ============================================================ */

.hero {
    padding: 28px 24px;
    border-radius: 18px;
    margin: 0 12px 18px 12px;
    background: linear-gradient(
        135deg,
        #17365d,
        #1f4e78
    );
    color: white;
    box-shadow: 0 8px 25px rgba(
        0,
        0,
        0,
        0.12
    );
    text-align: center;
}

.hero-inner {
    display: flex;
    align-items: center;
    justify-content: center;
    gap: 16px;
    flex-wrap: wrap;
}

.hero-text {
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
}

#hero-logo {
    width: 96px !important;
    height: 96px !important;
    min-width: 96px !important;
    min-height: 96px !important;
    border-radius: 50% !important;
    object-fit: contain !important;
    box-shadow: 0 10px 20px rgba(0, 0, 0, 0.18) !important;
    background: rgba(255, 255, 255, 0.12) !important;
    padding: 6px !important;
    display: block !important;
    margin: 0 auto !important;
}

#hero-logo img {
    width: 100% !important;
    height: 100% !important;
    object-fit: contain !important;
    border-radius: 50% !important;
}

.hero h1 {
    font-size: 52px !important;
    margin: 0 0 14px 0;
    font-weight: 800 !important;
    line-height: 1.08 !important;
    letter-spacing: 0.015em;
    color: rgba(36, 40, 45, 0.96) !important;
    text-shadow: none !important;
}

.hero p {
    font-size: 24px !important;
    line-height: 1.5 !important;
    opacity: 0.95;
    margin: 6px 0;
    color: rgba(36, 40, 45, 0.96) !important;
}

/* ============================================================
   SECTION BOXES
   ============================================================ */

.section-box {
    border-radius: 15px;
    border: 1px solid #dce3eb;
    background: white;
    padding: 22px;
    margin-bottom: 18px;
    box-shadow: 0 4px 15px rgba(
        0,
        0,
        0,
        0.05
    );
}

/* ============================================================
   BUTTONS
   ============================================================ */

.gradio-container button {
    font-size: 18px !important;
    font-weight: 700 !important;
    line-height: 1.25 !important;
}

.primary-button {
    min-height: 58px !important;
    font-size: 20px !important;
    font-weight: 800 !important;
}

.secondary-button {
    min-height: 52px !important;
    font-size: 18px !important;
    font-weight: 700 !important;
}

.folder-button {
    min-height: 52px !important;
    font-size: 18px !important;
    font-weight: 700 !important;
}

/* ============================================================
   INPUTS / DROPDOWNS / SEARCH
   ============================================================ */

.gradio-container input,
.gradio-container textarea,
.gradio-container select {
    font-size: 18px !important;
    min-height: 48px !important;
}

.gradio-container textarea {
    line-height: 1.5 !important;
}

.gradio-container input::placeholder,
.gradio-container textarea::placeholder {
    font-size: 17px !important;
    opacity: 0.75 !important;
}

.gradio-container label {
    font-size: 18px !important;
    font-weight: 700 !important;
}

/* Dropdown option text */
.gradio-container [role="option"] {
    font-size: 18px !important;
}

/* ============================================================
   DATAFRAME / TABLE TEXT
   ============================================================ */

.gradio-container .table-wrap,
.gradio-container .table-wrap table,
.gradio-container .dataframe,
.gradio-container .dataframe table,
.gradio-container .dataframe td,
.gradio-container .dataframe th {
    font-size: 17px !important;
}

.gradio-container .dataframe th {
    font-weight: 800 !important;
}

/* ============================================================
   DASHBOARD / SUMMARY TABLES
   ============================================================ */

.gradio-container table {
    font-size: 17px !important;
}

.gradio-container table th {
    font-size: 17px !important;
    font-weight: 800 !important;
}

.gradio-container table td {
    font-size: 17px !important;
}

.dashboard-card {
    border-radius: 12px;
    padding: 12px;
}

/* ============================================================
   FILE / STATUS / LOG TEXT
   ============================================================ */

.gradio-container .file-preview,
.gradio-container .file-preview * {
    font-size: 17px !important;
}

.gradio-container code,
.gradio-container pre {
    font-size: 16px !important;
    line-height: 1.5 !important;
}

/* Generation log */
.gradio-container textarea[readonly] {
    font-size: 16px !important;
    line-height: 1.5 !important;
}

/* ============================================================
   FOOTER
   ============================================================ */

.footer {
    text-align: center;
    padding: 26px;
    color: #667085;
    font-size: 16px !important;
    line-height: 1.5 !important;
}

/* ============================================================
   STAFF MANAGEMENT ACCORDION
   ============================================================ */
.management-accordion {
    border: 2px solid #dce3eb !important;
    border-radius: 15px !important;
    margin-bottom: 18px !important;
    background: #ffffff !important;
    box-shadow: 0 4px 15px rgba(0,0,0,0.05) !important;
}

.management-accordion > .label-wrap {
    min-height: 64px !important;
    padding: 14px 18px !important;
    font-size: 22px !important;
    font-weight: 800 !important;
    cursor: pointer !important;
}

.management-accordion .icon {
    width: 26px !important;
    height: 26px !important;
}

.management-accordion .gradio-accordion-content {
    padding: 18px !important;
}

/* ============================================================
   RESPONSIVE SETTINGS
   ============================================================ */

@media (max-width: 900px) {

    .gradio-container {
        width: 98vw !important;
        padding-left: 12px !important;
        padding-right: 12px !important;
    }

    .gradio-container,
    .gradio-container p,
    .gradio-container span,
    .gradio-container label,
    .gradio-container input,
    .gradio-container textarea,
    .gradio-container button {
        font-size: 17px !important;
    }

    .hero h1 {
        font-size: 32px !important;
    }

    .hero p {
        font-size: 17px !important;
    }

    .gradio-container h2 {
        font-size: 27px !important;
    }

    .gradio-container h3 {
        font-size: 23px !important;
    }
}


/* ============================================================
   LARGE CLICKABLE COLLAPSED SECTION HEADERS
   ============================================================ */
.big-section-accordion > .label,
.big-section-accordion > button,
.big-section-accordion button,
.big-section-accordion > .label-wrap {
    font-size: 28px !important;
    line-height: 1.35 !important;
    font-weight: 900 !important;
    min-height: 68px !important;
    padding: 16px 20px !important;
}

.big-section-accordion svg {
    width: 28px !important;
    height: 28px !important;
}

.management-accordion > .label,
.management-accordion > button,
.management-accordion button,
.management-accordion > .label-wrap {
    font-size: 28px !important;
    line-height: 1.35 !important;
    font-weight: 900 !important;
    min-height: 72px !important;
    padding: 16px 20px !important;
}

.management-accordion svg {
    width: 30px !important;
    height: 30px !important;
}

"""

# ============================================================
# STARTUP
# ============================================================

print()
print("================================")
print("RKM DUTY ROSTER GUI")
print("================================")
print("STEP 7.6")
print("Starting Gradio interface...")
print()


# ============================================================
# BUILD GUI
# ============================================================

with gr.Blocks(
    title=APP_TITLE,
    fill_width=True,
    css=CUSTOM_CSS
) as demo:

    with gr.Group(elem_id="login-screen") as login_screen:
        with gr.Group(elem_id="login-card"):
            gr.HTML(
                """
                <div class="login-brand">
                    <h1>RKM Duty Roster</h1>
                    <p>Duty Management Portal</p>
                </div>
                """
            )

            login_username = gr.Textbox(
                label="Username",
                placeholder="Enter username",
                autofocus=True,
                elem_classes="login-input"
            )

            login_password = gr.Textbox(
                label="Password",
                placeholder="Enter password",
                type="password",
                elem_classes="login-input"
            )

            login_button = gr.Button(
                "Login",
                variant="primary",
                elem_classes="primary-button"
            )

            login_error = gr.Markdown(
                value="",
                elem_id="login-error"
            )

    logout_button = gr.Button(
        "↩ Logout",
        variant="stop",
        elem_id="logout-button",
        visible=False
    )

    # ========================================================
    # HERO
    # ========================================================

    logo_path = ensure_logo_asset()

    with gr.Row():
        if logo_path:
            gr.HTML(
                f"""
                <div id="hero-logo-wrap">
                    <img id="hero-logo" src="/logo.png" alt="RKM logo" />
                </div>
                """
            )
        else:
            gr.HTML('<div class="hero-logo-placeholder">🛡️</div>')

        gr.HTML(
            """
            <div class="hero">
                <div class="hero-inner">
                    <div class="hero-text">
                        <h1>RKM Duty Roster Generator</h1>
                        <p>
                            Automated • Fair • Leave-Aware • Professionally Formatted
                        </p>
                        <p>
                            Step 7.6 — Staff & Department Management
                        </p>
                    </div>
                </div>
            </div>
            """
        )

    # ========================================================
    # FILE STATUS
    # ========================================================

    with gr.Accordion(
        "📁 INPUT & OUTPUT FILE STATUS    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 📁 Input & Output File Status"
        )

        file_status = gr.Markdown(
            value=get_file_status()
        )

        with gr.Row():

            check_files_button = gr.Button(
                "🔎 Check Input Files",
                variant="secondary",
                elem_classes="secondary-button"
            )

            refresh_files_button = gr.Button(
                "🔄 Refresh Status",
                variant="secondary",
                elem_classes="secondary-button"
            )

            open_input_button = gr.Button(
                "📂 Open Input Folder",
                variant="secondary",
                elem_classes="folder-button"
            )

            open_output_button = gr.Button(
                "📂 Open Output Folder",
                variant="secondary",
                elem_classes="folder-button"
            )

        input_validation_message = gr.Markdown(
            value=(
                "Click **Check Input Files** "
                "to validate the Excel files."
            )
        )

        input_validation_table = gr.Dataframe(
            headers=[
                "File",
                "Status",
                "Details",
                "Rows"
            ],
            datatype=[
                "str",
                "str",
                "str",
                "number"
            ],
            value=pd.DataFrame(),
            interactive=False,
            wrap=True
        )

    # ========================================================
    # STEP 7.6 - STAFF & DEPARTMENT MANAGEMENT
    # COLLAPSIBLE: CLICK THE HEADER TO OPEN/CLOSE
    # ========================================================

    with gr.Accordion(
        "👥 STAFF & DEPARTMENT MANAGEMENT    ▼",
        open=False,
        elem_classes="management-accordion big-section-accordion"
    ):

        gr.Markdown(
            """
### 👥 Staff & Department Management

Add, view, edit and remove employees and departments/duties **without manually editing the master Excel files**.

**Important:** Close `staff_master.xlsx` and `duty_master.xlsx` in Excel before saving changes.
"""
        )

        master_status = gr.Markdown(
            value="### 🟢 Master management ready."
        )

        # --------------------------------------------------------
        # ADD NEW EMPLOYEE / DEPARTMENT
        # --------------------------------------------------------
        with gr.Row():
            with gr.Column():
                gr.Markdown("### ➕ Add New Employee")

                employee_name_input = gr.Textbox(
                    label="Staff Name",
                    placeholder="Enter full employee name..."
                )

                employee_category_input = gr.Textbox(
                    label="Category",
                    placeholder="Example: Security / Staff / Front Desk..."
                )

                employee_active_input = gr.Dropdown(
                    label="Active",
                    choices=["YES", "NO"],
                    value="YES",
                    allow_custom_value=False
                )

                employee_duties_input = gr.CheckboxGroup(
                    label="Duty Permissions",
                    choices=get_employee_duty_choices(),
                    value=[],
                    info="Select the duties this employee is permitted to perform. Morning / Afternoon / Hostel Afternoon / Hostel Night / Night are supported."
                )

                add_employee_button = gr.Button(
                    "➕ ADD NEW EMPLOYEE",
                    variant="primary",
                    elem_classes="primary-button"
                )

            with gr.Column():
                gr.Markdown("### ➕ Add New Department / Duty")

                duty_name_input = gr.Textbox(
                    label="Department / Duty Name",
                    placeholder="Example: Pharmacy"
                )

                required_staff_input = gr.Number(
                    label="Required Staff",
                    value=1,
                    precision=0,
                    minimum=1,
                    step=1
                )

                add_duty_button = gr.Button(
                    "➕ ADD NEW DEPARTMENT / DUTY",
                    variant="primary",
                    elem_classes="primary-button"
                )

                refresh_master_button = gr.Button(
                    "🔄 REFRESH MASTER DATA",
                    variant="secondary",
                    elem_classes="secondary-button"
                )

                # --------------------------------------------------------
                # DUTY MASTER EXCEL UPLOAD - PROMINENT / ALWAYS VISIBLE
                # --------------------------------------------------------
                gr.Markdown("## 📤 UPLOAD DUTY MASTER EXCEL")
                gr.Markdown(
                    "Select your existing Duty Excel file. It must contain **Duty** and **Required Staff** columns. "
                    "The current duty_master.xlsx is backed up automatically before replacement."
                )
                duty_master_upload = gr.File(
                    label="🏥 Select Duty Master Excel",
                    file_types=[".xlsx", ".xls"],
                    type="filepath"
                )
                upload_duty_master_button = gr.Button(
                    "📤 UPLOAD & REPLACE DUTY MASTER",
                    variant="primary",
                    elem_classes="primary-button"
                )

                gr.Markdown("### ⚡ Quick Add Standard RKM Duties")
                standard_duties_input = gr.CheckboxGroup(
                    label="Select standard duties",
                    choices=STANDARD_RKM_DUTIES,
                    value=[],
                    info="Adds the selected duty to Duty Master and creates the matching Staff Master permission column."
                )
                add_standard_duties_button = gr.Button(
                    "⚡ ADD SELECTED STANDARD DUTIES",
                    variant="secondary",
                    elem_classes="secondary-button"
                )

        with gr.Row():
            staff_master_preview = gr.Dataframe(
                label="👥 Current Staff Master",
                value=get_master_tables()[0],
                interactive=False,
                wrap=True
            )

            duty_master_preview = gr.Dataframe(
                label="🏥 Current Duty Master",
                value=get_master_tables()[1],
                interactive=False,
                wrap=True
            )

        # --------------------------------------------------------
        # EMPLOYEE VIEW / EDIT / DELETE
        # --------------------------------------------------------
        gr.Markdown("## 👤 Employee View / Edit / Remove")

        with gr.Row():
            employee_selector = gr.Dropdown(
                label="Select Employee",
                choices=_staff_names_from_df(get_master_tables()[0]),
                value=None,
                allow_custom_value=False
            )
            view_employee_button = gr.Button(
                "👁️ VIEW",
                variant="secondary",
                elem_classes="secondary-button"
            )

        employee_view = gr.Markdown(
            "### ⚪ Select an employee and click **VIEW**."
        )

        # Hidden/visible original name tracker: this keeps the selected
        # record stable even if the user edits the employee name.
        employee_original_name = gr.Textbox(
            label="Original Employee Name",
            visible=False
        )

        with gr.Row():
            edit_employee_name = gr.Textbox(label="Employee Name")
            edit_employee_category = gr.Textbox(label="Category")
            edit_employee_active = gr.Dropdown(
                label="Active",
                choices=["YES", "NO"],
                value="YES",
                allow_custom_value=False
            )

        edit_employee_duties = gr.CheckboxGroup(
            label="Duty Permissions",
            choices=get_employee_duty_choices(),
            value=[]
        )

        with gr.Row():
            edit_employee_button = gr.Button(
                "💾 SAVE EMPLOYEE EDIT",
                variant="primary",
                elem_classes="primary-button"
            )
            confirm_employee_delete = gr.Checkbox(
                label="I have viewed the details and confirm DELETE",
                value=False
            )
            delete_employee_button = gr.Button(
                "🗑️ DELETE EMPLOYEE",
                variant="stop"
            )

        # --------------------------------------------------------
        # DEPARTMENT VIEW / EDIT / DELETE
        # --------------------------------------------------------
        gr.Markdown("## 🏥 Department / Duty View / Edit / Remove")

        with gr.Row():
            department_selector = gr.Dropdown(
                label="Select Department / Duty",
                choices=get_master_duty_choices(),
                value=None,
                allow_custom_value=False
            )
            view_department_button = gr.Button(
                "👁️ VIEW",
                variant="secondary",
                elem_classes="secondary-button"
            )

        department_view = gr.Markdown(
            "### ⚪ Select a department/duty and click **VIEW**."
        )

        department_original_name = gr.Textbox(
            label="Original Department / Duty Name",
            visible=False
        )

        with gr.Row():
            edit_department_name = gr.Textbox(label="Department / Duty Name")
            edit_department_required_staff = gr.Number(
                label="Required Staff",
                value=1,
                precision=0,
                minimum=1,
                step=1
            )

        with gr.Row():
            edit_department_button = gr.Button(
                "💾 SAVE DEPARTMENT / DUTY EDIT",
                variant="primary",
                elem_classes="primary-button"
            )
            confirm_department_delete = gr.Checkbox(
                label="I have viewed the details and confirm DELETE",
                value=False
            )
            delete_department_button = gr.Button(
                "🗑️ DELETE DEPARTMENT / DUTY",
                variant="stop"
            )

    # ========================================================
    # GENERATE ROSTER
    # ========================================================

    with gr.Group(
        elem_classes="section-box"
    ):

        gr.Markdown(
            "## 🚀 Generate Duty Roster"
        )

        gr.Markdown(
            """
The GUI uses your existing **app.py** roster engine.

Your Excel input files remain the source of the roster.

**Important:** Please do not click Generate repeatedly.
The system automatically prevents double-generation.
"""
        )

        with gr.Row():

            generate_button = gr.Button(
                "🚀 GENERATE DUTY ROSTER",
                variant="primary",
                elem_classes="primary-button"
            )

            refresh_output_button = gr.Button(
                "🔄 LOAD EXISTING OUTPUT",
                variant="secondary",
                elem_classes="secondary-button"
            )

            refresh_everything_button = gr.Button(
                "🔄 REFRESH EVERYTHING",
                variant="secondary",
                elem_classes="secondary-button"
            )

        generation_status = gr.Markdown(
            value=get_generation_status()
        )

    # ========================================================
    # DASHBOARD
    # ========================================================

    with gr.Accordion(
        "📊 ROSTER DASHBOARD    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 📊 Roster Dashboard"
        )

        dashboard_output = gr.Markdown(
            value=create_dashboard(
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame()
            )
        )

    # ========================================================
    # SUMMARY
    # ========================================================

    with gr.Accordion(
        "📋 GENERATION SUMMARY    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 📋 Generation Summary"
        )

        summary_output = gr.Markdown(
            value=create_summary(
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame()
            )
        )

    # ========================================================
    # FILTER & SEARCH
    # ========================================================

    with gr.Group(
        elem_classes="section-box"
    ):

        gr.Markdown(
            "## 🔎 Roster Filter & Search Dashboard"
        )

        gr.Markdown(
            """
Select **Date**, **Staff**, **Duty**, or use the
**Search** box. All filters always work from the
complete original roster.
"""
        )

        with gr.Row():

            date_filter = gr.Dropdown(
                choices=["ALL"],
                value="ALL",
                label="📅 Date",
                allow_custom_value=False
            )

            staff_filter = gr.Dropdown(
                choices=["ALL"],
                value="ALL",
                label="👤 Staff",
                allow_custom_value=False
            )

            duty_filter = gr.Dropdown(
                choices=["ALL"],
                value="ALL",
                label="🛡️ Duty",
                allow_custom_value=False
            )

        search_box = gr.Textbox(
            label="🔍 Search",
            placeholder=(
                "Search staff, duty, date, "
                "assignment type or reason..."
            ),
            lines=1
        )

        with gr.Row():

            apply_filter_button_component = gr.Button(
                "🔎 APPLY FILTER",
                variant="primary",
                elem_classes="secondary-button"
            )

            reset_filter_button = gr.Button(
                "↩️ RESET FILTERS",
                variant="secondary",
                elem_classes="secondary-button"
            )

        filter_status = gr.Markdown(
            value="### ⚪ No roster data available."
        )

    # ========================================================
    # ROSTER PREVIEW
    # ========================================================

    with gr.Accordion(
        "👀 ROSTER PREVIEW    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 👀 Roster Preview"
        )

        roster_preview = gr.Dataframe(
            value=pd.DataFrame(),
            label="Filtered Roster",
            interactive=False,
            wrap=True
        )

    # ========================================================
    # FILTERED EXCEL EXPORT
    # ========================================================

    with gr.Group(
        elem_classes="section-box"
    ):

        gr.Markdown(
            "## 📥 Filtered Excel Export"
        )

        gr.Markdown(
            """
Export only the records currently matching your
Date / Staff / Duty / Search filters.
"""
        )

        with gr.Row():

            export_filtered_button = gr.Button(
                "📥 EXPORT FILTERED EXCEL",
                variant="primary",
                elem_classes="secondary-button"
            )

            export_all_button = gr.Button(
                "📗 DOWNLOAD COMPLETE ROSTER",
                variant="secondary",
                elem_classes="secondary-button"
            )

        filtered_export_status = gr.Markdown(
            value=""
        )

        filtered_download_file = gr.File(
            label="📥 Filtered Excel",
            interactive=False
        )

    # ========================================================
    # STAFF SUMMARY
    # ========================================================

    with gr.Accordion(
        "👥 STAFF DUTY SUMMARY    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 👥 Staff Duty Summary"
        )

        staff_summary_preview = gr.Dataframe(
            value=pd.DataFrame(),
            interactive=False,
            wrap=True
        )

    # ========================================================
    # DUTY SUMMARY
    # ========================================================

    with gr.Accordion(
        "🛡️ DUTY SUMMARY    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 🛡️ Duty Summary"
        )

        duty_summary_preview = gr.Dataframe(
            value=pd.DataFrame(),
            interactive=False,
            wrap=True
        )

    # ========================================================
    # VALIDATION
    # ========================================================

    with gr.Accordion(
        "✅ ROSTER VALIDATION    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## ✅ Roster Validation"
        )

        validation_preview = gr.Dataframe(
            value=pd.DataFrame(),
            interactive=False,
            wrap=True
        )

    # ========================================================
    # COVER REPORT
    # ========================================================

    with gr.Accordion(
        "🟡 COVER REPORT    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 🟡 Cover Report"
        )

        cover_preview = gr.Dataframe(
            value=pd.DataFrame(),
            interactive=False,
            wrap=True
        )

    # ========================================================
    # COMPLETE EXCEL OUTPUT
    # ========================================================

    with gr.Group(
        elem_classes="section-box"
    ):

        gr.Markdown(
            "## 📗 Complete Excel Output"
        )

        gr.Markdown(
            """
Download the complete professionally formatted
`duty_roster.xlsx` workbook.
"""
        )

        download_file = gr.File(
            label="📗 Complete Duty Roster Excel",
            interactive=False
        )

    # ========================================================
    # GENERATION LOG
    # ========================================================

    with gr.Accordion(
        "📝 DETAILED GENERATION LOG    ▼",
        open=False,
        elem_classes="section-box big-section-accordion"
    ):

        gr.Markdown(
            "## 📝 Detailed Generation Log"
        )

        generation_log = gr.Textbox(
            label="Console / Generation Log",
            lines=25,
            max_lines=40,
            interactive=False
        )

    # ========================================================
    # ACTIVE RULES
    # ========================================================

    with gr.Group(
        elem_classes="section-box"
    ):

        gr.Markdown(
            RULES_MARKDOWN
        )

    # ========================================================
    # FOOTER
    # ========================================================

    gr.HTML(
        """
        <div class="footer">
            RKM Duty Roster Generator • Step 7.6<br>
            Automated Duty Scheduling • Staff Management • Department Management • Filter • Search • Export
        </div>
        """
    )

    # ========================================================
    # EVENTS
    # ========================================================

    login_button.click(
        fn=login_user,
        inputs=[login_username, login_password],
        outputs=[login_screen, login_error, logout_button]
    )

    login_password.submit(
        fn=login_user,
        inputs=[login_username, login_password],
        outputs=[login_screen, login_error, logout_button]
    )

    logout_button.click(
        fn=logout_user,
        inputs=[],
        outputs=[login_screen, logout_button]
    )

    # --------------------------------------------------------
    # CHECK INPUT FILES
    # --------------------------------------------------------

    check_files_button.click(
        fn=check_input_files,
        inputs=[],
        outputs=[
            input_validation_message,
            input_validation_table
        ]
    )

    # --------------------------------------------------------
    # REFRESH FILE STATUS
    # --------------------------------------------------------

    refresh_files_button.click(
        fn=get_file_status,
        inputs=[],
        outputs=[
            file_status
        ]
    )

    # --------------------------------------------------------
    # OPEN INPUT FOLDER
    # --------------------------------------------------------

    open_input_button.click(
        fn=open_input_folder,
        inputs=[],
        outputs=[
            file_status
        ]
    )

    # --------------------------------------------------------
    # OPEN OUTPUT FOLDER
    # --------------------------------------------------------

    open_output_button.click(
        fn=open_output_folder,
        inputs=[],
        outputs=[
            file_status
        ]
    )

    # --------------------------------------------------------
    # ADD NEW EMPLOYEE
    # --------------------------------------------------------

    add_employee_button.click(
        fn=add_new_employee,
        inputs=[
            employee_name_input,
            employee_category_input,
            employee_active_input,
            employee_duties_input
        ],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            employee_duties_input
        ]
    ).then(
        fn=refresh_master_data,
        inputs=[],
        outputs=[
            staff_master_preview,
            duty_master_preview,
            employee_duties_input,
            master_status
        ]
    ).then(
        fn=lambda: gr.update(choices=_staff_names_from_df(_get_staff_and_duty_tables()[0]), value=None),
        inputs=[],
        outputs=[employee_selector]
    ).then(
        fn=lambda: gr.update(choices=get_employee_duty_choices(), value=[]),
        inputs=[],
        outputs=[edit_employee_duties]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # ADD NEW DEPARTMENT / DUTY
    # --------------------------------------------------------

    add_duty_button.click(
        fn=add_new_department,
        inputs=[
            duty_name_input,
            required_staff_input
        ],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            employee_duties_input
        ]
    ).then(
        fn=refresh_master_data,
        inputs=[],
        outputs=[
            staff_master_preview,
            duty_master_preview,
            employee_duties_input,
            master_status
        ]
    ).then(
        fn=lambda: gr.update(choices=_staff_names_from_df(_get_staff_and_duty_tables()[0]), value=None),
        inputs=[],
        outputs=[employee_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=None),
        inputs=[],
        outputs=[department_selector]
    ).then(
        fn=lambda: gr.update(choices=get_employee_duty_choices(), value=[]),
        inputs=[],
        outputs=[edit_employee_duties]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # QUICK ADD STANDARD RKM DUTIES
    # --------------------------------------------------------

    add_standard_duties_button.click(
        fn=add_standard_rkm_duties,
        inputs=[standard_duties_input],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            employee_duties_input
        ]
    ).then(
        fn=lambda: gr.update(choices=_staff_names_from_df(_get_staff_and_duty_tables()[0]), value=None),
        inputs=[],
        outputs=[employee_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=None),
        inputs=[],
        outputs=[department_selector]
    ).then(
        fn=lambda: gr.update(choices=get_employee_duty_choices(), value=[]),
        inputs=[],
        outputs=[edit_employee_duties]
    ).then(
        fn=lambda: gr.update(value=[]),
        inputs=[],
        outputs=[standard_duties_input]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # UPLOAD DUTY MASTER EXCEL
    # --------------------------------------------------------

    upload_duty_master_button.click(
        fn=upload_duty_master_excel,
        inputs=[duty_master_upload],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            employee_duties_input
        ]
    ).then(
        fn=lambda: gr.update(choices=_staff_names_from_df(_get_staff_and_duty_tables()[0]), value=None),
        inputs=[],
        outputs=[employee_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=None),
        inputs=[],
        outputs=[department_selector]
    ).then(
        fn=lambda: gr.update(choices=get_employee_duty_choices(), value=[]),
        inputs=[],
        outputs=[edit_employee_duties]
    ).then(
        fn=lambda: gr.update(value=None),
        inputs=[],
        outputs=[duty_master_upload]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # REFRESH MASTER DATA
    # --------------------------------------------------------

    refresh_master_button.click(
        fn=refresh_master_data,
        inputs=[],
        outputs=[
            staff_master_preview,
            duty_master_preview,
            employee_duties_input,
            master_status
        ]
    ).then(
        fn=lambda: gr.update(choices=_staff_names_from_df(_get_staff_and_duty_tables()[0]), value=None),
        inputs=[],
        outputs=[employee_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=None),
        inputs=[],
        outputs=[department_selector]
    ).then(
        fn=lambda: gr.update(choices=get_employee_duty_choices(), value=[]),
        inputs=[],
        outputs=[edit_employee_duties]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # VIEW EMPLOYEE
    # --------------------------------------------------------

    view_employee_button.click(
        fn=view_employee_details,
        inputs=[employee_selector],
        outputs=[
            employee_view,
            employee_selector,
            edit_employee_name,
            edit_employee_category,
            edit_employee_active,
            edit_employee_duties,
            confirm_employee_delete
        ]
    ).then(
        fn=lambda name: name,
        inputs=[employee_selector],
        outputs=[employee_original_name]
    )

    # --------------------------------------------------------
    # SAVE EMPLOYEE EDIT
    # --------------------------------------------------------

    edit_employee_button.click(
        fn=edit_employee_details,
        inputs=[
            employee_original_name,
            edit_employee_name,
            edit_employee_category,
            edit_employee_active,
            edit_employee_duties
        ],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            employee_selector,
            department_selector,
            employee_original_name,
            edit_employee_name,
            edit_employee_category,
            edit_employee_active,
            edit_employee_duties
        ]
    ).then(
        fn=lambda: gr.update(choices=_staff_names_from_df(_get_staff_and_duty_tables()[0]), value=None),
        inputs=[],
        outputs=[employee_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=None),
        inputs=[],
        outputs=[department_selector]
    ).then(
        fn=lambda: False,
        inputs=[],
        outputs=[confirm_employee_delete]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # DELETE EMPLOYEE
    # --------------------------------------------------------

    delete_employee_button.click(
        fn=delete_employee_details,
        inputs=[
            employee_selector,
            confirm_employee_delete
        ],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            employee_selector,
            department_selector,
            employee_original_name,
            edit_employee_name,
            edit_employee_category,
            edit_employee_active,
            edit_employee_duties,
            confirm_employee_delete
        ]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # VIEW DEPARTMENT / DUTY
    # --------------------------------------------------------

    view_department_button.click(
        fn=view_department_details,
        inputs=[department_selector],
        outputs=[
            department_view,
            department_selector,
            edit_department_name,
            edit_department_required_staff
        ]
    ).then(
        fn=lambda name: name,
        inputs=[department_selector],
        outputs=[department_original_name]
    )

    # --------------------------------------------------------
    # SAVE DEPARTMENT / DUTY EDIT
    # --------------------------------------------------------

    edit_department_button.click(
        fn=edit_department_details,
        inputs=[
            department_original_name,
            edit_department_name,
            edit_department_required_staff
        ],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            department_selector,
            employee_duties_input,
            department_original_name,
            edit_department_name,
            edit_department_required_staff
        ]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=None),
        inputs=[],
        outputs=[department_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=[]),
        inputs=[],
        outputs=[employee_duties_input]
    ).then(
        fn=lambda: gr.update(choices=get_employee_duty_choices(), value=[]),
        inputs=[],
        outputs=[edit_employee_duties]
    ).then(
        fn=lambda: False,
        inputs=[],
        outputs=[confirm_department_delete]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # DELETE DEPARTMENT / DUTY
    # --------------------------------------------------------

    delete_department_button.click(
        fn=delete_department_details,
        inputs=[
            department_selector,
            confirm_department_delete
        ],
        outputs=[
            master_status,
            staff_master_preview,
            duty_master_preview,
            department_selector,
            employee_duties_input,
            department_original_name,
            edit_department_name,
            edit_department_required_staff,
            confirm_department_delete
        ]
    ).then(
        fn=lambda: gr.update(choices=_staff_names_from_df(_get_staff_and_duty_tables()[0]), value=None),
        inputs=[],
        outputs=[employee_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=None),
        inputs=[],
        outputs=[department_selector]
    ).then(
        fn=lambda: gr.update(choices=get_master_duty_choices(), value=[]),
        inputs=[],
        outputs=[employee_duties_input]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[file_status]
    )

    # --------------------------------------------------------
    # GENERATE
    # --------------------------------------------------------

    # IMPORTANT:
    # generate_roster() returns EXACTLY 14 values.

    generate_button.click(
        fn=generate_roster,
        inputs=[],
        outputs=[
            generation_status,
            dashboard_output,
            summary_output,
            roster_preview,
            staff_summary_preview,
            duty_summary_preview,
            validation_preview,
            cover_preview,
            generation_log,
            download_file,
            date_filter,
            staff_filter,
            duty_filter,
            filter_status
        ]
    ).then(
        fn=get_file_status,
        inputs=[],
        outputs=[
            file_status
        ]
    )

    # --------------------------------------------------------
    # LOAD EXISTING OUTPUT
    # --------------------------------------------------------

    refresh_output_button.click(
        fn=refresh_output,
        inputs=[],
        outputs=[
            generation_status,
            dashboard_output,
            summary_output,
            roster_preview,
            staff_summary_preview,
            duty_summary_preview,
            validation_preview,
            cover_preview,
            date_filter,
            staff_filter,
            duty_filter,
            roster_preview,
            filter_status,
            download_file
        ]
    )

    # --------------------------------------------------------
    # REFRESH EVERYTHING
    # --------------------------------------------------------

    refresh_everything_button.click(
        fn=refresh_everything,
        inputs=[],
        outputs=[
            file_status,
            input_validation_message,
            input_validation_table,
            generation_status,
            dashboard_output,
            summary_output,
            roster_preview,
            staff_summary_preview,
            duty_summary_preview,
            validation_preview,
            cover_preview,
            date_filter,
            staff_filter,
            duty_filter,
            filter_status,
            download_file
        ]
    )

    # --------------------------------------------------------
    # APPLY FILTER
    # --------------------------------------------------------

    apply_filter_button_component.click(
        fn=apply_filter_button,
        inputs=[
            date_filter,
            staff_filter,
            duty_filter,
            search_box
        ],
        outputs=[
            roster_preview,
            filter_status
        ]
    )

    # --------------------------------------------------------
    # DATE CHANGE
    # --------------------------------------------------------

    date_filter.change(
        fn=apply_filter_button,
        inputs=[
            date_filter,
            staff_filter,
            duty_filter,
            search_box
        ],
        outputs=[
            roster_preview,
            filter_status
        ]
    )

    # --------------------------------------------------------
    # STAFF CHANGE
    # --------------------------------------------------------

    staff_filter.change(
        fn=apply_filter_button,
        inputs=[
            date_filter,
            staff_filter,
            duty_filter,
            search_box
        ],
        outputs=[
            roster_preview,
            filter_status
        ]
    )

    # --------------------------------------------------------
    # DUTY CHANGE
    # --------------------------------------------------------

    duty_filter.change(
        fn=apply_filter_button,
        inputs=[
            date_filter,
            staff_filter,
            duty_filter,
            search_box
        ],
        outputs=[
            roster_preview,
            filter_status
        ]
    )

    # --------------------------------------------------------
    # SEARCH ENTER
    # --------------------------------------------------------

    search_box.submit(
        fn=apply_filter_button,
        inputs=[
            date_filter,
            staff_filter,
            duty_filter,
            search_box
        ],
        outputs=[
            roster_preview,
            filter_status
        ]
    )

    # --------------------------------------------------------
    # RESET FILTER
    # --------------------------------------------------------

    reset_filter_button.click(
        fn=reset_filters,
        inputs=[],
        outputs=[
            date_filter,
            staff_filter,
            search_box,
            duty_filter,
            roster_preview,
            filter_status
        ]
    )

    # --------------------------------------------------------
    # EXPORT FILTERED
    # --------------------------------------------------------

    export_filtered_button.click(
        fn=export_filtered_excel,
        inputs=[
            date_filter,
            staff_filter,
            duty_filter,
            search_box
        ],
        outputs=[
            filtered_download_file,
            filtered_export_status
        ]
    )

    # --------------------------------------------------------
    # DOWNLOAD COMPLETE ROSTER
    # --------------------------------------------------------

    def get_complete_output_file():

        if OUTPUT_FILE.exists():

            return str(
                OUTPUT_FILE
            )

        return None

    export_all_button.click(
        fn=get_complete_output_file,
        inputs=[],
        outputs=[
            download_file
        ]
    )

    # --------------------------------------------------------
    # AUTO LOAD ON START
    # --------------------------------------------------------

    demo.load(
        fn=refresh_everything,
        inputs=[],
        outputs=[
            file_status,
            input_validation_message,
            input_validation_table,
            generation_status,
            dashboard_output,
            summary_output,
            roster_preview,
            staff_summary_preview,
            duty_summary_preview,
            validation_preview,
            cover_preview,
            date_filter,
            staff_filter,
            duty_filter,
            filter_status,
            download_file
        ]
    )


    # --------------------------------------------------------
    # AUTO LOAD MASTER DATA ON START
    # --------------------------------------------------------

    demo.load(
        fn=refresh_master_data,
        inputs=[],
        outputs=[
            staff_master_preview,
            duty_master_preview,
            employee_duties_input,
            master_status
        ]
    )

# ============================================================
# LAUNCH
# ============================================================

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 7860))
    demo.launch(
        server_name="0.0.0.0",
        server_port=7861,
        inbrowser=False,
<<<<<<< HEAD
        share=True
    )
=======
    )
>>>>>>> 59f80a7607990e89c63b7bb0bd6491097089dc9c
