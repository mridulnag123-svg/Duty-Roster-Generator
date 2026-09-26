import os
import sys
import subprocess
import traceback
import tkinter as tk
from tkinter import messagebox

import pandas as pd
from pymongo import MongoClient

from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    Alignment,
    Border,
    Side,
    PatternFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


# ============================================================
# MONGODB ATLAS CLOUD DATABASE CONNECTION
# ============================================================
try:
    client = MongoClient(
        "mongodb+srv://mridulnag123_db_user:Nl45jEl0uLUeqPs@cluster0.o2ekegg.mongodb.net/?retryWrites=true&w=majority&appName=Cluster0"
    )
    db = client["duty_roster_db"]
    roster_collection = db["roster_records"]
    print("Connected to MongoDB Atlas successfully.")
except Exception as e:
    print("MongoDB Connection Error:", e)


# ============================================================
# SHARED EXCEL BORDER STYLE
# ============================================================
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


# ============================================================
# RKM DUTY ROSTER GENERATOR
# STEP 7 - GUI + STEP 6 ROSTER ENGINE
# ============================================================


# ============================================================
# FILE PATHS
# ============================================================

availability_file = "data/input/availability.xlsx"
staff_master_file = "data/master/staff_master.xlsx"
duty_master_file = "data/master/duty_master.xlsx"

output_file = "data/output/duty_roster.xlsx"

os.makedirs(
    "data/output",
    exist_ok=True
)


# ============================================================
# SETTINGS
# ============================================================

NORMAL_MAX_DUTIES_PER_DAY = 1

SHORTAGE_MAX_DUTIES_PER_DAY = 2

SECOND_DUTY_ASSIGNMENT_TYPE = "COVER"

AUTO_AVAILABLE_ACTIVE_STAFF = True

MAX_CONSECUTIVE_NIGHT_DUTIES = 3

# ============================================================
# DYNAMIC DUTY STRUCTURE BY AVAILABLE RKMs
# ============================================================
DYNAMIC_DUTY_STRUCTURE_ENABLED = True
DUTY_STRUCTURE_BY_STAFF_COUNT = {
    3: ["Morning", "Afternoon", "Night", "Hostel Night"],
    4: ["Morning", "Afternoon", "Hostel Night", "Night"],
    5: ["Morning", "Afternoon", "Hostel Afternoon", "Hostel Night", "Night"],
}


# ============================================================
# GLOBAL
# ============================================================

last_generated_file = output_file


# ============================================================
# GENERATE ROSTER
# ============================================================

def generate_roster():

    global last_generated_file

    print("\n================================")
    print("LOADING INPUT FILES")
    print("================================")

    try:

        availability_df = pd.read_excel(
            availability_file
        )

        staff_df = pd.read_excel(
            staff_master_file
        )

        duty_df = pd.read_excel(
            duty_master_file
        )

        print(
            "Input files loaded successfully."
        )

    except Exception as e:

        print("\n================================")
        print("ERROR")
        print("================================")

        print(e)

        raise


    # ========================================================
    # CLEAN COLUMN NAMES
    # ========================================================

    availability_df.columns = (
        availability_df.columns
        .astype(str)
        .str.strip()
    )

    staff_df.columns = (
        staff_df.columns
        .astype(str)
        .str.strip()
    )

    duty_df.columns = (
        duty_df.columns
        .astype(str)
        .str.strip()
    )


    # ========================================================
    # REQUIRED COLUMNS
    # ========================================================

    required_availability_columns = [
        "Date",
        "Available",
        "Leave"
    ]

    required_staff_columns = [
        "Staff Name",
        "Active"
    ]

    required_duty_columns = [
        "Duty",
        "Required Staff"
    ]


    for column in required_availability_columns:

        if column not in availability_df.columns:

            raise ValueError(
                f"Availability file missing column: {column}"
            )


    for column in required_staff_columns:

        if column not in staff_df.columns:

            raise ValueError(
                f"Staff master missing column: {column}"
            )


    for column in required_duty_columns:

        if column not in duty_df.columns:

            raise ValueError(
                f"Duty master missing column: {column}"
            )


    # ========================================================
    # CLEAN AVAILABILITY
    # ========================================================

    availability_df = availability_df[
        availability_df["Date"]
        .astype(str)
        .str.strip()
        .str.lower()
        != "date"
    ].copy()


    availability_df["Date"] = pd.to_datetime(
        availability_df["Date"],
        errors="coerce"
    )


    availability_df = availability_df.dropna(
        subset=["Date"]
    ).copy()


    availability_df = (
        availability_df
        .sort_values("Date")
        .reset_index(drop=True)
    )


    if availability_df.empty:

        raise ValueError(
            "No valid dates found in availability.xlsx."
        )


    # ========================================================
    # CLEAN STAFF MASTER
    # ========================================================

    staff_df["Staff Name"] = (
        staff_df["Staff Name"]
        .astype(str)
        .str.strip()
    )

    staff_df["Active"] = (
        staff_df["Active"]
        .astype(str)
        .str.strip()
        .str.upper()
    )


    # ========================================================
    # REMOVE INVALID STAFF
    # ========================================================

    staff_df = staff_df[
        staff_df["Staff Name"].notna()
        &
        (
            staff_df["Staff Name"]
            .astype(str)
            .str.strip()
            != ""
        )
    ].copy()


    # ========================================================
    # DUPLICATE STAFF CHECK
    # ========================================================

    staff_df["_Staff Name Key"] = (
        staff_df["Staff Name"]
        .astype(str)
        .str.strip()
        .str.upper()
    )


    duplicate_staff_names = (
        staff_df["_Staff Name Key"]
        .duplicated(keep=False)
    )


    if duplicate_staff_names.any():

        duplicate_names = (
            staff_df.loc[
                duplicate_staff_names,
                "Staff Name"
            ]
            .drop_duplicates()
            .tolist()
        )

        raise ValueError(
            "Duplicate staff names found:\n\n"
            + "\n".join(
                f"- {name}"
                for name in duplicate_names
            )
        )


    staff_df = staff_df.drop(
        columns=["_Staff Name Key"]
    )


    # ========================================================
    # ACTIVE STAFF
    # ========================================================

    active_staff_df = staff_df[
        staff_df["Active"] == "YES"
    ].copy()


    if active_staff_df.empty:

        raise ValueError(
            "No active staff found."
        )


    staff_names = (
        active_staff_df["Staff Name"]
        .tolist()
    )


    print(
        f"Active staff loaded: {len(staff_names)}"
    )


    for staff_name in staff_names:

        print(
            f"  - {staff_name}"
        )


    # ========================================================
    # DUTY LIST
    # ========================================================

    duties = []


    for _, row in duty_df.iterrows():

        duty_name = str(
            row["Duty"]
        ).strip()

        if not duty_name:
            continue

        try:

            required_staff = int(
                row["Required Staff"]
            )

        except Exception:

            required_staff = 1


        if required_staff < 1:
            continue


        for _ in range(required_staff):

            duties.append(
                duty_name
            )


    if not duties:

        raise ValueError(
            "No duties found in duty_master.xlsx."
        )


    # ========================================================
    # DUTY PERMISSION WARNING
    # ========================================================

    missing_permission_columns = []


    for duty_name in sorted(
        set(duties)
    ):

        if duty_name not in active_staff_df.columns:

            missing_permission_columns.append(
                duty_name
            )


    if missing_permission_columns:

        print("\n================================")
        print("WARNING")
        print("================================")

        print(
            "Missing duty permission columns:"
        )

        for duty_name in missing_permission_columns:

            print(
                f" - {duty_name}"
            )


    # ========================================================
    # COUNTERS
    # ========================================================

    duty_count = {
        staff_name: 0
        for staff_name in staff_names
    }


    primary_count = {
        staff_name: 0
        for staff_name in staff_names
    }


    cover_count = {
        staff_name: 0
        for staff_name in staff_names
    }


    staff_duty_history = {
        staff_name: []
        for staff_name in staff_names
    }


    last_assignment_date = {
        staff_name: None
        for staff_name in staff_names
    }


    last_duty_by_staff = {
        staff_name: None
        for staff_name in staff_names
    }


    last_cover_date = {
        staff_name: None
        for staff_name in staff_names
    }


    # ========================================================
    # NIGHT HISTORY
    # ========================================================

    night_history = {
        staff_name: []
        for staff_name in staff_names
    }


    # ========================================================
    # OUTPUT DATA
    # ========================================================

    roster = []

    leave_records = []

    daily_duty_records = {}


    # ========================================================
    # HELPERS
    # ========================================================

    def get_staff_info(staff_name):

        result = active_staff_df[
            active_staff_df["Staff Name"]
            .astype(str)
            .str.upper()
            ==
            str(staff_name)
            .strip()
            .upper()
        ]

        if result.empty:

            return None

        return result.iloc[0]


    def is_allowed_for_duty(
        staff_name,
        duty_name
    ):

        staff_info = get_staff_info(
            staff_name
        )

        if staff_info is None:

            return False


        if duty_name in active_staff_df.columns:

            permission = str(
                staff_info[duty_name]
            ).strip().upper()

            return permission == "YES"


        return True


    def is_night_duty(duty_name):

        return "NIGHT" in str(
            duty_name
        ).strip().upper()


    def get_daily_duty_slots(eligible_staff):
        available_count = len(eligible_staff)

        if not DYNAMIC_DUTY_STRUCTURE_ENABLED:
            return duties

        configured = DUTY_STRUCTURE_BY_STAFF_COUNT.get(available_count)
        if not configured:
            return duties

        master_map = {}
        for master_duty in duties:
            key = str(master_duty).strip().casefold()
            if key not in master_map:
                master_map[key] = str(master_duty).strip()

        selected = []
        missing = []
        for configured_duty in configured:
            key = str(configured_duty).strip().casefold()
            if key in master_map:
                selected.append(master_map[key])
            else:
                missing.append(configured_duty)

        if missing:
            print(
                f"WARNING: {available_count} RKM duty structure "
                f"needs missing Duty Master duty(s): {', '.join(missing)}"
            )
            print("Using Duty Master duties for this date instead.")
            return duties

        return selected


    def has_same_duty_consecutive(
        staff_name,
        duty_name
    ):

        history = staff_duty_history.get(
            staff_name,
            []
        )

        if not history:

            return False

        return history[-1] == duty_name


    def has_three_consecutive_night(
        staff_name,
        duty_date
    ):

        history = night_history.get(
            staff_name,
            []
        )

        if len(history) < MAX_CONSECUTIVE_NIGHT_DUTIES:

            return False


        recent_dates = history[
            -MAX_CONSECUTIVE_NIGHT_DUTIES:
        ]


        for i in range(
            1,
            len(recent_dates)
        ):

            if (
                recent_dates[i]
                - recent_dates[i - 1]
            ).days != 1:

                return False


        if (
            duty_date
            - recent_dates[-1]
        ).days != 1:

            return False


        return True


    def can_receive_duty_today(
        staff_name,
        duty_name,
        daily_count,
        allow_second_duty,
        duty_date
    ):

        maximum = (
            SHORTAGE_MAX_DUTIES_PER_DAY
            if allow_second_duty
            else NORMAL_MAX_DUTIES_PER_DAY
        )


        if daily_count.get(
            staff_name,
            0
        ) >= maximum:

            return False


        today_duties = (
            daily_duty_records
            .get(
                staff_name,
                []
            )
        )


        if duty_name in today_duties:

            return False


        if has_same_duty_consecutive(
            staff_name,
            duty_name
        ):

            return False


        if is_night_duty(
            duty_name
        ):

            if has_three_consecutive_night(
                staff_name,
                duty_date
            ):

                return False


        return True


    # ========================================================
    # PRIMARY SELECTION
    # ========================================================

    def select_best_primary_staff(
        candidates,
        duty_name,
        duty_date
    ):

        if not candidates:

            return None


        def score(staff_name):

            total_count = duty_count.get(
                staff_name,
                0
            )

            staff_primary_count = (
                primary_count.get(
                    staff_name,
                    0
                )
            )

            staff_cover_count = (
                cover_count.get(
                    staff_name,
                    0
                )
            )

            same_duty_count = (
                staff_duty_history
                .get(
                    staff_name,
                    []
                )
                .count(
                    duty_name
                )
            )

            last_date = (
                last_assignment_date.get(
                    staff_name
                )
            )


            if last_date is None:

                days_since_last = 999999

            else:

                days_since_last = (
                    duty_date - last_date
                ).days


            return (
                total_count,
                staff_primary_count,
                staff_cover_count,
                same_duty_count,
                -days_since_last,
                staff_name.upper()
            )


        return min(
            candidates,
            key=score
        )


    # ========================================================
    # COVER SELECTION
    # ========================================================

    def select_best_cover_staff(
        candidates,
        duty_name,
        duty_date
    ):

        if not candidates:

            return None


        def score(staff_name):

            total_count = duty_count.get(
                staff_name,
                0
            )

            staff_cover_count = (
                cover_count.get(
                    staff_name,
                    0
                )
            )

            staff_primary_count = (
                primary_count.get(
                    staff_name,
                    0
                )
            )

            last_cover = (
                last_cover_date.get(
                    staff_name
                )
            )


            if last_cover is None:

                days_since_cover = 999999

            else:

                days_since_cover = (
                    duty_date - last_cover
                ).days


            same_duty_count = (
                staff_duty_history
                .get(
                    staff_name,
                    []
                )
                .count(
                    duty_name
                )
            )


            today_count = len(
                daily_duty_records.get(
                    staff_name,
                    []
                )
            )


            return (
                total_count,
                staff_cover_count,
                staff_primary_count,
                today_count,
                same_duty_count,
                -days_since_cover,
                staff_name.upper()
            )


        return min(
            candidates,
            key=score
        )


    # ========================================================
    # PROCESS DATES
    # ========================================================

    for _, availability_row in availability_df.iterrows():

        duty_date = availability_row["Date"]

        date_text = duty_date.strftime(
            "%d/%m/%Y"
        )


        daily_duty_records = {
            staff_name: []
            for staff_name in staff_names
        }


        # ====================================================
        # LEAVE
        # ====================================================

        leave_text = str(
            availability_row["Leave"]
        ).strip()


        leave_staff = []


        if leave_text.lower() not in [
            "",
            "-",
            "nan",
            "none"
        ]:

            for name in leave_text.split(","):

                name = name.strip()

                if not name:
                    continue


                staff_info = get_staff_info(
                    name
                )


                if staff_info is None:

                    print(
                        f"WARNING: Leave staff '{name}' "
                        f"not found in Staff Master."
                    )

                    continue


                official_name = (
                    staff_info["Staff Name"]
                )


                if official_name not in leave_staff:

                    leave_staff.append(
                        official_name
                    )


                leave_records.append({

                    "Date":
                        date_text,

                    "Staff":
                        official_name,

                    "Status":
                        "LEAVE"

                })


        # ====================================================
        # ACTIVE STAFF AUTOMATICALLY AVAILABLE
        # ====================================================

        if AUTO_AVAILABLE_ACTIVE_STAFF:

            eligible_staff = [

                staff_name

                for staff_name in staff_names

                if staff_name not in leave_staff

            ]

        else:

            available_text = str(
                availability_row["Available"]
            ).strip()


            available_staff = []


            if available_text.lower() not in [
                "",
                "nan",
                "none"
            ]:

                for name in available_text.split(","):

                    name = name.strip()

                    if not name:
                        continue


                    staff_info = get_staff_info(
                        name
                    )


                    if staff_info is None:

                        continue


                    official_name = (
                        staff_info["Staff Name"]
                    )


                    if official_name not in available_staff:

                        available_staff.append(
                            official_name
                        )


            eligible_staff = [

                name

                for name in available_staff

                if name not in leave_staff

            ]


        # ====================================================
        # PRINT LEAVE
        # ====================================================

        for staff_name in leave_staff:

            print(
                f"LEAVE: {staff_name} skipped on "
                f"{date_text}"
            )


        primary_assigned_today = set()


        # ====================================================
        # DYNAMIC DUTY STRUCTURE FOR THIS DATE
        # ====================================================
        daily_duties = get_daily_duty_slots(eligible_staff)

        print(
            f"DUTY STRUCTURE: {len(eligible_staff)} available RKM(s) -> "
            f"{', '.join(daily_duties)}"
        )

        # ====================================================
        # DUTIES
        # ====================================================

        for duty_name in daily_duties:


            # =================================================
            # PRIMARY CANDIDATES
            # =================================================

            primary_candidates = []


            daily_counts = {

                name: len(
                    daily_duty_records.get(
                        name,
                        []
                    )
                )

                for name in staff_names

            }


            for staff_name in eligible_staff:

                if staff_name in primary_assigned_today:

                    continue


                if not is_allowed_for_duty(
                    staff_name,
                    duty_name
                ):

                    continue


                if not can_receive_duty_today(
                    staff_name,
                    duty_name,
                    daily_counts,
                    False,
                    duty_date
                ):

                    continue


                primary_candidates.append(
                    staff_name
                )


            # =================================================
            # PRIMARY ASSIGNMENT
            # =================================================

            if primary_candidates:

                selected_staff = (
                    select_best_primary_staff(
                        primary_candidates,
                        duty_name,
                        duty_date
                    )
                )


                roster.append({

                    "Date":
                        date_text,

                    "Duty":
                        duty_name,

                    "Staff":
                        selected_staff,

                    "Assignment Type":
                        "PRIMARY",

                    "Reason":
                        "Normal assignment"

                })


                duty_count[
                    selected_staff
                ] += 1


                primary_count[
                    selected_staff
                ] += 1


                staff_duty_history[
                    selected_staff
                ].append(
                    duty_name
                )


                last_duty_by_staff[
                    selected_staff
                ] = duty_name


                last_assignment_date[
                    selected_staff
                ] = duty_date


                daily_duty_records[
                    selected_staff
                ].append(
                    duty_name
                )


                if is_night_duty(
                    duty_name
                ):

                    night_history[
                        selected_staff
                    ].append(
                        duty_date
                    )


                primary_assigned_today.add(
                    selected_staff
                )


                continue


            # =================================================
            # SECOND DUTY / COVER
            # =================================================

            second_duty_candidates = []


            daily_counts = {

                name: len(
                    daily_duty_records.get(
                        name,
                        []
                    )
                )

                for name in staff_names

            }


            for staff_name in eligible_staff:

                if not is_allowed_for_duty(
                    staff_name,
                    duty_name
                ):

                    continue


                if not can_receive_duty_today(
                    staff_name,
                    duty_name,
                    daily_counts,
                    True,
                    duty_date
                ):

                    continue


                second_duty_candidates.append(
                    staff_name
                )


            if second_duty_candidates:

                selected_staff = (
                    select_best_cover_staff(
                        second_duty_candidates,
                        duty_name,
                        duty_date
                    )
                )


                roster.append({

                    "Date":
                        date_text,

                    "Duty":
                        duty_name,

                    "Staff":
                        selected_staff,

                    "Assignment Type":
                        SECOND_DUTY_ASSIGNMENT_TYPE,

                    "Reason":
                        "Staff shortage - second duty"

                })


                duty_count[
                    selected_staff
                ] += 1


                cover_count[
                    selected_staff
                ] += 1


                staff_duty_history[
                    selected_staff
                ].append(
                    duty_name
                )


                last_duty_by_staff[
                    selected_staff
                ] = duty_name


                last_assignment_date[
                    selected_staff
                ] = duty_date


                last_cover_date[
                    selected_staff
                ] = duty_date


                daily_duty_records[
                    selected_staff
                ].append(
                    duty_name
                )


                if is_night_duty(
                    duty_name
                ):

                    night_history[
                        selected_staff
                    ].append(
                        duty_date
                    )


                print(
                    f"COVER: {selected_staff} assigned to "
                    f"{duty_name} on "
                    f"{date_text} "
                    f"(SECOND DUTY - STAFF SHORTAGE)"
                )


            # =================================================
            # UNASSIGNED
            # =================================================

            else:

                roster.append({

                    "Date":
                        date_text,

                    "Duty":
                        duty_name,

                    "Staff":
                        "UNASSIGNED",

                    "Assignment Type":
                        "UNASSIGNED",

                    "Reason":
                        "No eligible staff available"

                })


                print(
                    f"WARNING: {duty_name} UNASSIGNED on "
                    f"{date_text}"
                )


    # ========================================================
    # ROSTER DATAFRAME
    # ========================================================

    roster_df = pd.DataFrame(
        roster
    )


    # ========================================================
    # VALIDATION
    # ========================================================

    validation_results = []


    # ========================================================
    # CHECK 1 - UNASSIGNED
    # ========================================================

    unassigned_count = len(
        roster_df[
            roster_df["Assignment Type"]
            == "UNASSIGNED"
        ]
    )


    validation_results.append({

        "Check":
            "Unassigned duties",

        "Status":
            "PASS"
            if unassigned_count == 0
            else "WARNING",

        "Details":
            "All duties have been assigned."
            if unassigned_count == 0
            else
            f"{unassigned_count} duty/duties unassigned."

    })


    # ========================================================
    # ASSIGNED DATAFRAME
    # ========================================================

    assigned_df = roster_df[
        roster_df["Staff"] != "UNASSIGNED"
    ].copy()


    # ========================================================
    # CHECK 2 - MAX TWO DUTIES
    # ========================================================

    if not assigned_df.empty:

        daily_staff_counts = (
            assigned_df
            .groupby(
                ["Date", "Staff"]
            )
            .size()
        )

    else:

        daily_staff_counts = pd.Series(
            dtype=int
        )


    violation_count = (
        daily_staff_counts
        > SHORTAGE_MAX_DUTIES_PER_DAY
    ).sum()


    validation_results.append({

        "Check":
            "Maximum two duties per staff per day",

        "Status":
            "PASS"
            if violation_count == 0
            else "FAIL",

        "Details":
            "No staff has more than two duties on the same day."
            if violation_count == 0
            else
            f"{violation_count} violation(s)."

    })


    # ========================================================
    # CHECK 3 - LEAVE
    # ========================================================

    leave_violation_count = 0


    for _, row in roster_df.iterrows():

        if row["Staff"] == "UNASSIGNED":

            continue


        for record in leave_records:

            if (
                record["Date"] == row["Date"]
                and
                record["Staff"] == row["Staff"]
            ):

                leave_violation_count += 1


    validation_results.append({

        "Check":
            "Leave staff not assigned",

        "Status":
            "PASS"
            if leave_violation_count == 0
            else "FAIL",

        "Details":
            "No leave staff assigned."
            if leave_violation_count == 0
            else
            f"{leave_violation_count} violation(s)."

    })


    # ========================================================
    # CHECK 4 - SAME DUTY CONSECUTIVE
    # ========================================================

    same_duty_consecutive = []


    if not assigned_df.empty:

        check_df = assigned_df.copy()

        check_df["DateObject"] = pd.to_datetime(
            check_df["Date"],
            format="%d/%m/%Y",
            errors="coerce"
        )


        check_df = check_df.sort_values(
            ["Staff", "DateObject"]
        )


        for staff_name, staff_group in check_df.groupby(
            "Staff"
        ):

            previous_date = None

            previous_duty = None


            for _, row in staff_group.iterrows():

                current_date = row["DateObject"]

                current_duty = row["Duty"]


                if (
                    previous_date is not None
                    and
                    previous_duty == current_duty
                    and
                    (
                        current_date
                        - previous_date
                    ).days == 1
                ):

                    same_duty_consecutive.append({

                        "Staff":
                            staff_name,

                        "Duty":
                            current_duty,

                        "Previous Date":
                            previous_date.strftime(
                                "%d/%m/%Y"
                            ),

                        "Current Date":
                            current_date.strftime(
                                "%d/%m/%Y"
                            )

                    })


                previous_date = current_date

                previous_duty = current_duty


    validation_results.append({

        "Check":
            "Same duty on consecutive days",

        "Status":
            "PASS"
            if not same_duty_consecutive
            else "FAIL",

        "Details":
            "No same duty assigned on consecutive days."
            if not same_duty_consecutive
            else
            f"{len(same_duty_consecutive)} violation(s)."

    })


    # ========================================================
    # CHECK 5 - SAME DUTY SAME DAY
    # ========================================================

    same_day_same_duty = []


    if not assigned_df.empty:

        counts = (
            assigned_df
            .groupby(
                ["Date", "Staff", "Duty"]
            )
            .size()
        )


        for index, count in counts.items():

            if count > 1:

                same_day_same_duty.append({

                    "Date":
                        index[0],

                    "Staff":
                        index[1],

                    "Duty":
                        index[2],

                    "Count":
                        int(count)

                })


    validation_results.append({

        "Check":
            "Same duty twice on same day",

        "Status":
            "PASS"
            if not same_day_same_duty
            else "FAIL",

        "Details":
            "No duplicate same-day duty found."
            if not same_day_same_duty
            else
            f"{len(same_day_same_duty)} violation(s)."

    })


    # ========================================================
    # CHECK 6 - COVER DISTRIBUTION
    # ========================================================

    cover_values = list(
        cover_count.values()
    )


    cover_difference = (
        max(cover_values)
        - min(cover_values)
        if cover_values
        else 0
    )


    validation_results.append({

        "Check":
            "Cover distribution",

        "Status":
            "PASS"
            if cover_difference <= 2
            else "WARNING",

        "Details":
            "Cover duties are reasonably balanced."
            if cover_difference <= 2
            else
            "Cover duties are unevenly distributed."

    })


    # ========================================================
    # CHECK 7 - DUTY COVERAGE
    # ========================================================

    total_expected_duties = (
        len(availability_df)
        *
        len(duties)
    )


    total_generated = len(
        roster_df
    )


    validation_results.append({

        "Check":
            "Duty coverage",

        "Status":
            "PASS"
            if total_generated == total_expected_duties
            else "WARNING",

        "Details":
            f"{total_generated} assignments generated."
            if total_generated == total_expected_duties
            else
            f"Expected {total_expected_duties}, "
            f"generated {total_generated}."

    })


    # ========================================================
    # CHECK 8 - ACTIVE STAFF DISTRIBUTION
    # ========================================================

    zero_duty_staff = [

        name

        for name in staff_names

        if duty_count.get(
            name,
            0
        ) == 0

    ]


    validation_results.append({

        "Check":
            "Active staff distribution",

        "Status":
            "PASS"
            if not zero_duty_staff
            else "WARNING",

        "Details":
            "Every active staff member received at least one duty."
            if not zero_duty_staff
            else
            "Active staff with zero duties: "
            + ", ".join(zero_duty_staff)

    })


    # ========================================================
    # CHECK 9 - PERMISSION
    # ========================================================

    permission_violations = []


    for _, row in assigned_df.iterrows():

        if not is_allowed_for_duty(
            row["Staff"],
            row["Duty"]
        ):

            permission_violations.append({

                "Date":
                    row["Date"],

                "Staff":
                    row["Staff"],

                "Duty":
                    row["Duty"]

            })


    validation_results.append({

        "Check":
            "Duty permission",

        "Status":
            "PASS"
            if not permission_violations
            else "FAIL",

        "Details":
            "All assigned staff have permission for their duties."
            if not permission_violations
            else
            f"{len(permission_violations)} permission violation(s)."

    })


    # ========================================================
    # CHECK 10 - NIGHT DUTIES
    # ========================================================

    night_consecutive_violations = []


    if not assigned_df.empty:

        night_df = assigned_df[
            assigned_df["Duty"]
            .astype(str)
            .str.upper()
            .str.contains(
                "NIGHT",
                na=False
            )
        ].copy()


        if not night_df.empty:

            night_df["DateObject"] = pd.to_datetime(
                night_df["Date"],
                format="%d/%m/%Y",
                errors="coerce"
            )


            for staff_name, group in night_df.groupby(
                "Staff"
            ):

                group = group.sort_values(
                    "DateObject"
                )


                consecutive = 1

                previous_date = None


                for _, row in group.iterrows():

                    current_date = row["DateObject"]


                    if previous_date is not None:

                        if (
                            current_date
                            - previous_date
                        ).days == 1:

                            consecutive += 1

                        else:

                            consecutive = 1


                    if consecutive > MAX_CONSECUTIVE_NIGHT_DUTIES:

                        night_consecutive_violations.append({

                            "Staff":
                                staff_name,

                            "Date":
                                row["Date"],

                            "Duty":
                                row["Duty"],

                            "Consecutive NIGHT Duties":
                                consecutive

                        })


                    previous_date = current_date


    validation_results.append({

        "Check":
            "Maximum 3 consecutive NIGHT duties",

        "Status":
            "PASS"
            if not night_consecutive_violations
            else "FAIL",

        "Details":
            "No staff has more than 3 consecutive NIGHT duties."
            if not night_consecutive_violations
            else
            f"{len(night_consecutive_violations)} violation(s)."

    })


    # ========================================================
    # VALIDATION DATAFRAME
    # ========================================================

    validation_df = pd.DataFrame(
        validation_results
    )


    # ========================================================
    # REPORTS
    # ========================================================

    staff_summary_df = pd.DataFrame({

        "Staff":
            staff_names,

        "Total Duties":
            [
                duty_count[name]
                for name in staff_names
            ],

        "Primary Duties":
            [
                primary_count[name]
                for name in staff_names
            ],

        "Cover Duties":
            [
                cover_count[name]
                for name in staff_names
            ]

    })


    staff_summary_df = (
        staff_summary_df
        .sort_values(
            by=[
                "Total Duties",
                "Primary Duties",
                "Cover Duties"
            ],
            ascending=[
                False,
                False,
                False
            ]
        )
        .reset_index(drop=True)
    )


    duty_summary_df = (

        roster_df[
            roster_df["Staff"] != "UNASSIGNED"
        ]

        .groupby("Duty")
        .size()

        .reset_index(
            name="Assignments"
        )

    )


    leave_df = pd.DataFrame(
        leave_records
    )


    if leave_df.empty:

        leave_df = pd.DataFrame(
            columns=[
                "Date",
                "Staff",
                "Status"
            ]
        )


    cover_df = roster_df[
        roster_df["Assignment Type"]
        == "COVER"
    ].copy()


    unassigned_df = roster_df[
        roster_df["Assignment Type"]
        == "UNASSIGNED"
    ].copy()


    cover_summary = pd.DataFrame({

        "Staff":
            staff_names,

        "Cover Duties":
            [
                cover_count[name]
                for name in staff_names
            ]

    })


    cover_summary = (
        cover_summary
        .sort_values(
            "Cover Duties",
            ascending=False
        )
        .reset_index(drop=True)
    )


    consecutive_df = pd.DataFrame(
        same_duty_consecutive
    )


    if consecutive_df.empty:

        consecutive_df = pd.DataFrame(
            columns=[
                "Staff",
                "Duty",
                "Previous Date",
                "Current Date"
            ]
        )


    same_day_duplicate_df = pd.DataFrame(
        same_day_same_duty
    )


    if same_day_duplicate_df.empty:

        same_day_duplicate_df = pd.DataFrame(
            columns=[
                "Date",
                "Staff",
                "Duty",
                "Count"
            ]
        )


    permission_violation_df = pd.DataFrame(
        permission_violations
    )


    if permission_violation_df.empty:

        permission_violation_df = pd.DataFrame(
            columns=[
                "Date",
                "Staff",
                "Duty"
            ]
        )


    night_consecutive_df = pd.DataFrame(
        night_consecutive_violations
    )


    if night_consecutive_df.empty:

        night_consecutive_df = pd.DataFrame(
            columns=[
                "Staff",
                "Date",
                "Duty",
                "Consecutive NIGHT Duties"
            ]
        )


    # ========================================================
    # ROSTER MATRIX
    # ========================================================

    if not roster_df.empty:

        roster_matrix = (

            roster_df

            .pivot_table(

                index="Date",

                columns="Duty",

                values="Staff",

                aggfunc="first"

            )

            .reset_index()

        )

    else:

        roster_matrix = pd.DataFrame()


    # ========================================================
    # PROFESSIONAL DUTY LIST / STAFF MATRIX
    # ========================================================
    duty_list_records = []

    if not roster_df.empty:

        roster_for_matrix = roster_df.copy()

        roster_for_matrix["_Date"] = pd.to_datetime(
            roster_for_matrix["Date"],
            dayfirst=True,
            errors="coerce"
        )

        roster_for_matrix = roster_for_matrix.dropna(
            subset=["_Date"]
        )

        for duty_date in sorted(
            roster_for_matrix["_Date"].drop_duplicates().tolist()
        ):

            row = {
                "Date": duty_date
            }

            day_rows = roster_for_matrix[
                roster_for_matrix["_Date"] == duty_date
            ]

            for staff_name in staff_names:

                staff_rows = day_rows[
                    day_rows["Staff"].astype(str).str.strip().str.upper()
                    == str(staff_name).strip().upper()
                ]

                assigned = []

                for _, assignment_row in staff_rows.iterrows():

                    duty_name = str(
                        assignment_row.get("Duty", "")
                    ).strip()

                    assignment_type = str(
                        assignment_row.get("Assignment Type", "")
                    ).strip().upper()

                    if not duty_name:
                        continue

                    if assignment_type == "COVER":
                        duty_display = f"{duty_name} (COVER)"
                    elif assignment_type == "UNASSIGNED":
                        duty_display = "UNASSIGNED"
                    else:
                        duty_display = duty_name

                    if duty_display not in assigned:
                        assigned.append(duty_display)

                row[staff_name] = " + ".join(assigned) if assigned else "0"

            duty_list_records.append(row)

    else:
        for duty_date in pd.to_datetime(
            availability_df["Date"], errors="coerce"
        ).dropna().sort_values().tolist():
            row = {"Date": duty_date}
            for staff_name in staff_names:
                row[staff_name] = "0"
            duty_list_records.append(row)

    duty_list_df = pd.DataFrame(
        duty_list_records,
        columns=["Date"] + staff_names
    )

    # ========================================================
    # SAVE EXCEL & MONGODB
    # ========================================================

    print("\n================================")
    print("SAVING DATA TO EXCEL & MONGODB")
    print("================================")


    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        duty_list_df.to_excel(
            writer,
            sheet_name="Duty List",
            index=False,
            startrow=3
        )

        roster_df.to_excel(
            writer,
            sheet_name="Roster",
            index=False
        )

        roster_matrix.to_excel(
            writer,
            sheet_name="Roster Matrix",
            index=False
        )

        staff_summary_df.to_excel(
            writer,
            sheet_name="Staff Summary",
            index=False
        )

        duty_summary_df.to_excel(
            writer,
            sheet_name="Duty Summary",
            index=False
        )

        leave_df.to_excel(
            writer,
            sheet_name="Leave Report",
            index=False
        )

        cover_df.to_excel(
            writer,
            sheet_name="Cover Report",
            index=False
        )

        unassigned_df.to_excel(
            writer,
            sheet_name="Unassigned",
            index=False
        )

        active_staff_df.to_excel(
            writer,
            sheet_name="Active Staff",
            index=False
        )

        cover_summary.to_excel(
            writer,
            sheet_name="Cover Summary",
            index=False
        )

        validation_df.to_excel(
            writer,
            sheet_name="Validation",
            index=False
        )

        consecutive_df.to_excel(
            writer,
            sheet_name="Consecutive Duty",
            index=False
        )

        same_day_duplicate_df.to_excel(
            writer,
            sheet_name="Same Day Duplicate",
            index=False
        )

        permission_violation_df.to_excel(
            writer,
            sheet_name="Permission Violations",
            index=False
        )

        night_consecutive_df.to_excel(
            writer,
            sheet_name="Night Consecutive",
            index=False
        )

    # Save Roster Records to MongoDB Cloud Database
    try:
        records_to_insert = roster_df.to_dict(orient="records")
        if records_to_insert:
            roster_collection.delete_many({}) # Clear old records
            roster_collection.insert_many(records_to_insert)
            print("Roster data successfully saved to MongoDB Atlas!")
    except Exception as db_err:
        print("Failed to save data to MongoDB:", db_err)


    # ========================================================
    # FORMAT EXCEL
    # ========================================================

    print("\n================================")
    print("FORMATTING EXCEL")
    print("================================")


    wb = load_workbook(
        output_file
    )


    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
        size=11
    )

    primary_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9"
    )

    cover_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )

    unassigned_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC"
    )

    pass_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9"
    )

    warning_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )

    fail_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC"
    )


    for ws in wb.worksheets:

        ws.freeze_panes = "A2"

        ws.sheet_view.showGridLines = False


        if ws.max_row >= 1:

            for cell in ws[1]:

                cell.font = header_font

                cell.fill = header_fill

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                cell.border = THIN_BORDER


        for row in ws.iter_rows():

            for cell in row:

                cell.border = THIN_BORDER

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True
                )


        for column_cells in ws.columns:

            max_length = 0

            column_letter = (
                get_column_letter(
                    column_cells[0].column
                )
            )


            for cell in column_cells:

                try:

                    length = len(
                        str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        length
                    )

                except Exception:

                    pass


            ws.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                45
            )


        ws.row_dimensions[1].height = 25

        ws.page_setup.orientation = "landscape"

        ws.page_setup.paperSize = (
            ws.PAPERSIZE_A4
        )

        ws.page_setup.fitToWidth = 1

        ws.page_setup.fitToHeight = 0

        ws.sheet_properties.pageSetUpPr.fitToPage = True

        ws.page_margins = PageMargins(
            left=0.25,
            right=0.25,
            top=0.5,
            bottom=0.5,
            header=0.2,
            footer=0.2
        )


    if "Duty List" in wb.sheetnames:

        ws = wb["Duty List"]

        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "B5"

        last_col = max(ws.max_column, 1)
        last_row = max(ws.max_row, 4)

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=last_col
        )
        ws.cell(1, 1).value = "Ramakrishna Mission Ashram, Narendrapur"
        ws.cell(1, 1).font = Font(
            name="Calibri",
            size=14,
            bold=True,
            color="1F4E78"
        )
        ws.cell(1, 1).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        ws.row_dimensions[1].height = 25

        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=last_col
        )
        ws.cell(2, 1).value = "RKM Duty List"
        ws.cell(2, 1).font = Font(
            name="Calibri",
            size=13,
            bold=True,
            color="000000"
        )
        ws.cell(2, 1).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        ws.row_dimensions[2].height = 23

        for cell in ws[4]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )
            cell.font = Font(
                name="Calibri",
                size=11,
                bold=True,
                color="FFFFFF"
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )
            cell.border = THIN_BORDER

        ws.row_dimensions[4].height = 28

        ws.column_dimensions["A"].width = 15

        for col in range(2, last_col + 1):
            ws.column_dimensions[
                get_column_letter(col)
            ].width = 18

        for row in range(5, last_row + 1):

            ws.cell(row, 1).number_format = "dd-mm-yyyy"
            ws.cell(row, 1).alignment = Alignment(
                horizontal="center",
                vertical="center"
            )
            ws.cell(row, 1).font = Font(
                name="Calibri",
                size=10,
                bold=True
            )
            ws.cell(row, 1).border = THIN_BORDER

            for col in range(2, last_col + 1):

                cell = ws.cell(row, col)
                value = str(cell.value or "").strip()
                upper_value = value.upper()

                cell.font = Font(
                    name="Calibri",
                    size=10,
                    bold=True if value not in ("", "0") else False
                )
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )
                cell.border = THIN_BORDER

                if value == "0" or value == "":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FFF2CC"
                    )
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=True,
                        color="7F6000"
                    )

                elif "UNASSIGNED" in upper_value:
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F4CCCC"
                    )
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=True,
                        color="C00000"
                    )

                elif "COVER" in upper_value:
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FFF2CC"
                    )
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=True,
                        color="C65911"
                    )

                elif "NIGHT" in upper_value:
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F4CCCC"
                    )
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=True,
                        color="C00000"
                    )

                else:
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="E2F0D9"
                    )
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=True,
                        color="1F4E78"
                    )

            ws.row_dimensions[row].height = 24

        ws.auto_filter.ref = (
            f"A4:{get_column_letter(last_col)}{last_row}"
        )

        ws.print_title_rows = "1:4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(
            left=0.25,
            right=0.25,
            top=0.4,
            bottom=0.4,
            header=0.2,
            footer=0.2
        )

        ws.oddFooter.center.text = "RKM Duty List"
        ws.oddFooter.center.size = 9


    if "Roster" in wb.sheetnames:

        ws = wb["Roster"]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        assignment_col = headers.get(
            "Assignment Type"
        )

        if assignment_col:

            for row in range(
                2,
                ws.max_row + 1
            ):

                assignment = ws.cell(
                    row=row,
                    column=assignment_col
                ).value

                if assignment == "PRIMARY":
                    fill = primary_fill
                elif assignment == "COVER":
                    fill = cover_fill
                elif assignment == "UNASSIGNED":
                    fill = unassigned_fill
                else:
                    fill = None

                if fill:
                    for cell in ws[row]:
                        cell.fill = fill


    if "Roster Matrix" in wb.sheetnames:

        ws = wb["Roster Matrix"]
        ws.freeze_panes = "B2"

        for row in ws.iter_rows(
            min_row=2
        ):
            for cell in row:
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )


    if "Validation" in wb.sheetnames:

        ws = wb["Validation"]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        status_col = headers.get(
            "Status"
        )

        if status_col:

            for row in range(
                2,
                ws.max_row + 1
            ):

                cell = ws.cell(
                    row=row,
                    column=status_col
                )

                if cell.value == "PASS":
                    cell.fill = pass_fill
                elif cell.value == "WARNING":
                    cell.fill = warning_fill
                elif cell.value == "FAIL":
                    cell.fill = fail_fill


    if "Night Consecutive" in wb.sheetnames:

        ws = wb[
            "Night Consecutive"
        ]

        ws.freeze_panes = "A2"

        for row in ws.iter_rows(
            min_row=2
        ):
            for cell in row:
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

        if ws.max_row > 1:
            for row in ws.iter_rows(
                min_row=2
            ):
                for cell in row:
                    cell.fill = fail_fill


    wb.save(
        output_file
    )

    last_generated_file = output_file


    print("\n================================")
    print("SUCCESS")
    print("================================")

    print("Duty roster saved successfully!")
    print("Output file:", output_file)
    print("Total assignments:", len(roster_df))

    return output_file


# ============================================================
# OPEN EXCEL
# ============================================================

def open_excel_file():

    try:

        if not os.path.exists(
            output_file
        ):

            messagebox.showwarning(
                "File Not Found",
                "Please generate the duty roster first."
            )

            return

        absolute_path = os.path.abspath(
            output_file
        )

        if sys.platform.startswith(
            "win"
        ):

            os.startfile(
                absolute_path
            )

        elif sys.platform == "darwin":

            subprocess.Popen([
                "open",
                absolute_path
            ])

        else:

            subprocess.Popen([
                "xdg-open",
                absolute_path
            ])

    except Exception as e:

        messagebox.showerror(
            "Error",
            f"Could not open Excel file.\n\n{e}"
        )


# ============================================================
# OPEN OUTPUT FOLDER
# ============================================================

def open_output_folder():

    try:

        folder = os.path.abspath(
            "data/output"
        )

        if sys.platform.startswith(
            "win"
        ):

            os.startfile(
                folder
            )

        elif sys.platform == "darwin":

            subprocess.Popen([
                "open",
                folder
            ])

        else:

            subprocess.Popen([
                "xdg-open",
                folder
            ])

    except Exception as e:

        messagebox.showerror(
            "Error",
            f"Could not open output folder.\n\n{e}"
        )


# ============================================================
# GUI
# ============================================================

root = tk.Tk()

root.title(
    "RKM Duty Roster Generator"
)

root.geometry(
    "720x560"
)

root.resizable(
    False,
    False
)

root.configure(
    bg="#F4F7FB"
)


# ============================================================
# TITLE
# ============================================================

title_label = tk.Label(

    root,

    text=
    "RKM DUTY ROSTER GENERATOR",

    font=(
        "Arial",
        22,
        "bold"
    ),

    fg="#1F4E78",

    bg="#F4F7FB"
)

title_label.pack(
    pady=(35, 5)
)


# ============================================================
# SUBTITLE
# ============================================================

subtitle_label = tk.Label(

    root,

    text=
    "Automatic Duty Roster Management System",

    font=(
        "Arial",
        12
    ),

    fg="#666666",

    bg="#F4F7FB"
)

subtitle_label.pack(
    pady=(0, 20)
)


# ============================================================
# RULES FRAME
# ============================================================

rules_frame = tk.Frame(

    root,

    bg="white",

    highlightbackground="#A6A6A6",

    highlightthickness=1
)

rules_frame.pack(
    padx=50,
    fill="x"
)


rules_title = tk.Label(

    rules_frame,

    text="System Rules Enabled",

    font=(
        "Arial",
        15,
        "bold"
    ),

    fg="#1F4E78",

    bg="white"
)

rules_title.pack(
    pady=(18, 10)
)


rules = [

    "✓ Active staff automatically available",

    "✓ Leave protection",

    "✓ Duty permission checking",

    "✓ Fair duty distribution",

    "✓ Maximum 2 duties per day during shortage",

    "✓ No same duty on consecutive days",

    "✓ Maximum 3 consecutive NIGHT duties",

    "✓ Automatic Excel validation & formatting"

]


for rule in rules:

    label = tk.Label(

        rules_frame,

        text=rule,

        font=(
            "Arial",
            11
        ),

        anchor="w",

        fg="#333333",

        bg="white"
    )

    label.pack(
        anchor="w",
        padx=150,
        pady=2
    )


# ============================================================
# STATUS LABEL
# ============================================================

status_label = tk.Label(

    root,

    text="Ready to generate roster.",

    font=(
        "Arial",
        10
    ),

    fg="#666666",

    bg="#F4F7FB"
)

status_label.pack(
    pady=(18, 8)
)


# ============================================================
# BUTTON FRAME
# ============================================================

button_frame = tk.Frame(

    root,

    bg="#F4F7FB"
)

button_frame.pack(
    pady=10
)


# ============================================================
# GENERATE FUNCTION
# ============================================================

def generate_from_gui():

    generate_button.config(
        state="disabled",
        text="GENERATING..."
    )


    status_label.config(
        text="Generating duty roster...",
        fg="#1F4E78"
    )


    root.update_idletasks()


    try:

        output = generate_roster()


        status_label.config(
            text="Roster generated successfully.",
            fg="#2E7D32"
        )


        result = messagebox.askyesno(

            "Success",

            "Duty roster generated successfully!\n\n"
            f"Output file:\n{output}\n\n"
            "Do you want to open the Excel file now?"

        )


        if result:

            open_excel_file()


    except Exception as e:

        print("\n================================")
        print("ERROR")
        print("================================")

        traceback.print_exc()


        status_label.config(
            text="Roster generation failed.",
            fg="#C00000"
        )


        messagebox.showerror(

            "Generation Error",

            "Duty roster could not be generated.\n\n"
            f"{e}"

        )


    finally:

        generate_button.config(
            state="normal",
            text="GENERATE DUTY ROSTER"
        )


# ============================================================
# GENERATE BUTTON
# ============================================================

generate_button = tk.Button(

    button_frame,

    text=
    "GENERATE DUTY ROSTER",

    command=
    generate_from_gui,

    font=(
        "Arial",
        14,
        "bold"
    ),

    fg="white",

    bg="#2E7D32",

    activebackground="#256628",

    activeforeground="white",

    width=27,

    height=2,

    relief="flat",

    cursor="hand2"
)

generate_button.pack(
    pady=5
)


# ============================================================
# OPEN EXCEL BUTTON
# ============================================================

open_excel_button = tk.Button(

    button_frame,

    text=
    "OPEN EXCEL FILE",

    command=
    open_excel_file,

    font=(
        "Arial",
        10,
        "bold"
    ),

    fg="#1F4E78",

    bg="white",

    activebackground="#EAF2F8",

    width=22,

    height=1,

    relief="solid",

    cursor="hand2"
)

open_excel_button.pack(
    pady=5
)


# ============================================================
# OPEN OUTPUT FOLDER BUTTON
# ============================================================

open_folder_button = tk.Button(

    button_frame,

    text=
    "OPEN OUTPUT FOLDER",

    command=
    open_output_folder,

    font=(
        "Arial",
        10,
        "bold"
    ),

    fg="#1F4E78",

    bg="white",

    activebackground="#EAF2F8",

    width=22,

    height=1,

    relief="solid",

    cursor="hand2"
)

open_folder_button.pack(
    pady=5
)


# ============================================================
# FOOTER
# ============================================================

footer_label = tk.Label(

    root,

    text=
    "RKM Duty Roster Generator • Step 7",

    font=(
        "Arial",
        9
    ),

    fg="#888888",

    bg="#F4F7FB"
)

footer_label.pack(
    side="bottom",
    pady=15
)


# ============================================================
# START GUI
# ============================================================

root.mainloop()
